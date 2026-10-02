"""Step 12b - fill the native Rating_Scorecard template for Godrej Properties Limited.

Historicals from compute_ratios.py, projections copied (never linked) from forward_model.py.

NORMALISATION DECISION, stated here rather than buried: PAT fed into the scorecard is REPORTED
PAT (total), NOT normalised for the Rs 2,093 cr (FY2026) / Rs 930 cr (FY2025) non-cash fair
value gain embedded in Other Income. This mirrors how ICRA's OWN published rationale for this
issuer treats it - ICRA flags the OPERATING MARGIN as modest "excluding fair value gain upon
acquisition/relinquishment of control" as a SEPARATE observation, while still using reported PAT
for headline profitability. OPBDIT (row 13) is unaffected either way - it excludes ALL Other
Income unconditionally under the ICRA convention (compute_ratios.py's opbdit()), so the fair
value gain never enters it regardless of this PAT choice.
"""

import os
import shutil
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import compute_ratios as C                                              # noqa: E402
import forward_model as F                                               # noqa: E402

SRC = "../../Rating_Scorecard.xlsx"
DST = "Rating_Scorecard_Godrej Properties Limited.xlsx"
HIST = ["FY2024", "FY2025", "FY2026"]

# Cash paid towards construction/land - proxied by (Collections - OCF), the same relationship
# used in forward_model.py, since a discrete "capex" cash-flow line is not meaningful for a
# real-estate developer (nearly all cash outflow IS the inventory/construction spend, which sits
# inside operating activities under Ind AS, not investing activities). FY2024 left blank - OCF
# and Collections are not disclosed as clean figures for FY2024 (see UNVERIFIED_INPUTS).
CONSTRUCTION_SPEND = {
    "FY2024": None,
    "FY2025": C.KPI["collections"][1] - C.KPI["ocf"][1],   # 17,047 - 7,484 = 9,563
    "FY2026": C.KPI["collections"][2] - C.KPI["ocf"][2],   # 19,965 - 7,830 = 12,135
}

# No dividend paid in any of the three years - Directors' Report, FY2025-26 AR (the company has
# historically retained earnings for land investment, consistent with its growth-stage profile).
DIVIDEND = {"FY2024": 0.0, "FY2025": 0.0, "FY2026": 0.0}

QUALITATIVE = {
    "business": (2, "India's #1 residential real estate developer by Booking Value for the "
                    "third consecutive year (company's own MD&A, corroborated across both "
                    "FY2025 and FY2026 annual reports), with a geographically diversified "
                    "portfolio across 10+ cities (MMR, NCR, Bengaluru, Pune, Chennai, "
                    "Ahmedabad, Kolkata, Hyderabad, Nagpur, Indore) and a 30+ year track "
                    "record. Booking Value grew from Rs 29,444 cr (FY2025) to Rs 34,171 cr "
                    "(FY2026), +16%, with Collections up 17% to Rs 19,965 cr - both cross-"
                    "checked across two independent annual reports. Not a 1, because revenue "
                    "recognition is concentrated on project completion timing (execution risk) "
                    "and reported OPBDIT margin is structurally negative-to-nil under the "
                    "project-completion method (see the engine's DISCLOSED_DIFFERENCES) - a "
                    "feature ICRA's own rationale independently flags as 'modest' operating "
                    "margins ex the fair-value gain."),
    "management": (2, "Unmodified/clean audit opinion on both standalone and consolidated "
                      "FY2026 financial statements ('give a true and fair view... in "
                      "conformity with the accounting principles generally accepted in "
                      "India') - Independent Auditor's Report, FY2025-26 AR. One technical "
                      "exception noted (Rule 11(g) reporting on accounting-software audit-"
                      "trail backup servers not located in India) - a systems/compliance "
                      "matter, not a financial-statement qualification or going-concern "
                      "flag. Part of the Godrej Industries Group with a long-tenured "
                      "leadership team (CEO with 9.5 years at GPL, 21 years in real estate)."),
    "support": (1, "ICRA's OWN published rationale (21-Jul-2026, reference library, same "
                   "issuer) explicitly credits 'exceptional financial flexibility as part of "
                   "the Godrej Industries Group', citing Rs 12,850 crore of equity raised "
                   "FY2019-FY2025 via private placement/QIP. This is a genuine, agency-"
                   "identified support factor - unlike a case with zero support, ONE notch is "
                   "entered here to reflect it, kept modest because the support is FINANCIAL "
                   "FLEXIBILITY (demonstrated capital-raising access) rather than an explicit "
                   "guarantee or a majority-shareholder cash-support undertaking."),
    "liquidity": ("Strong",
                  "Cash and cash equivalents of Rs 1,860.74 cr plus current investments of Rs "
                  "2,884.23 cr = Rs 4,744.97 cr (computed in compute_ratios.py, consolidated "
                  "BS, March 2026) - a narrower, more conservative basis than ICRA's own "
                  "stated Rs 8,970 cr 'free cash and liquid investments' figure (which "
                  "includes items such as bank balances/FDs beyond current investments and "
                  "the Rs 1,488 cr QIP proceeds this engine has not separately traced). "
                  "Operating cash flow of Rs 7,830 cr (FY2026, +5% YoY) comfortably covers "
                  "finance costs of Rs 136.86 cr - a cover of over 50 times. ICRA's own "
                  "rationale states the liquidity position 'would be strong even after "
                  "deployment' of its earmarked funds. Strong, not Superior, given the "
                  "company's own guidance of rising debt (to Rs 18,500-19,000 cr) in FY2027 "
                  "to fund land investment."),
    "outlook": ("Stable", "See the rating view in the workpaper."),
}


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                                     # noqa: BLE001
        pass

    shutil.copyfile(SRC, DST)
    wb = load_workbook(DST)
    ws = wb["Inputs"]
    R = C.compute()
    base = F.project(False)

    ws["C4"] = "Godrej Properties Limited"
    ws["C5"] = "ICRA"
    ws["C6"] = "Real Estate / Construction"    # exact label in Ref_Weights - "Real Estate" alone
                                                # does not match and leaves the anchor #N/A
    ws["C7"] = "02-Sep-2026"

    rows = {
        12: lambda y: R[y]["Revenue"],
        13: lambda y: R[y]["OPBDIT (ex Other Income)"],
        14: lambda y: C.PL["depreciation"][C.YEARS.index(y)],
        15: lambda y: C.PL["finance_cost"][C.YEARS.index(y)],
        16: lambda y: R[y]["PAT (total)"],                 # reported - see module docstring
        17: lambda y: R[y]["Total debt"],
        18: lambda y: R[y]["Cash + liquid investments"],
        19: lambda y: R[y]["Net worth"],
        20: lambda y: C.BS["total_liab"][C.YEARS.index(y)],
        21: lambda y: (C.BS["inventory"][C.YEARS.index(y)] + C.BS["trade_recv_c"][C.YEARS.index(y)]
                       + R[y]["Cash + liquid investments"]),
        22: lambda y: C.BS["borrow_c"][C.YEARS.index(y)],
        23: lambda y: CONSTRUCTION_SPEND[y],
        24: lambda y: DIVIDEND[y],
    }
    for r, fn in rows.items():
        for j, y in enumerate(HIST):
            v = fn(y)
            if v is not None:
                ws.cell(row=r, column=3 + j, value=round(v, 2))

    fy26_rev = R["FY2026"]["Revenue"]
    proj = {
        "revenue_growth": (round(base[0]["Collections"] / C.KPI["collections"][2] - 1, 4),
                           round(base[1]["Collections"] / base[0]["Collections"] - 1, 4)),
        # OPBDIT margin is not a meaningful projection driver here (see DISCLOSED_DIFFERENCES) -
        # held at the FY2026 actual rather than assuming an improvement the model cannot derive.
        "ebitda_margin": (R["FY2026"]["OPBDIT margin %"] / 100.0, R["FY2026"]["OPBDIT margin %"] / 100.0),
        "capex": (base[0]["Land/capex"], base[1]["Land/capex"]),
        "net_new_debt": (base[0]["New debt raised"], base[1]["New debt raised"]),
        "interest_rate": (0.0705, 0.0705),
        "tax_rate": (0.25, 0.25),
        "dividend_payout": (0.0, 0.0),
    }
    for r, key in [(29, "revenue_growth"), (30, "ebitda_margin"), (31, "capex"),
                   (32, "net_new_debt"), (33, "interest_rate"), (34, "tax_rate"),
                   (35, "dividend_payout")]:
        for j in range(2):
            ws.cell(row=r, column=3 + j, value=proj[key][j])

    ws["C39"] = None                                   # no override; Real Estate sector base
    ws["C40"] = QUALITATIVE["business"][0]
    ws["C41"] = QUALITATIVE["management"][0]
    ws["C42"] = QUALITATIVE["support"][0]
    ws["C43"] = QUALITATIVE["liquidity"][0]
    ws["C44"] = QUALITATIVE["outlook"][0]

    # ------------------------------------------------------- hard-cap disapplication (C67/D67)
    # The corporate hard cap "interest coverage < 1x -> force B+" tests ACCRUAL coverage. Under
    # Ind AS project-completion accounting a residential developer recognises revenue only on
    # occupancy certificate, so accrual OPBDIT is negative by construction in most years while
    # the business is strongly cash-generative. Firing the cap here moved the anchor from A+ to
    # B+ - nine of the twelve notches of divergence were this artefact alone.
    ws["C67"] = "N"
    ws["D67"] = (
        "DISAPPLIED. The cap tests accrual interest coverage (Ratios!E14 = -3.04x FY2026). "
        "Under Ind AS project-completion accounting GPL recognises revenue only on receipt of "
        "the occupancy certificate, so accrual OPBDIT is negative by construction in FY2024 "
        "and FY2026 and near-nil in FY2025 - the company's own MD&A states this creates 'a lag "
        "in reported earnings', and ICRA's own published rationale independently calls the "
        "reported operating margins 'modest'. The CASH coverage is the meaningful measure and "
        "is exceptionally strong: FY2026 operating cash flow of Rs 7,830 crore against finance "
        "costs of Rs 136.86 crore = 57.2x. Disapplying the cap moves the anchor from B+ "
        "(index 13) to A+ (index 4). This is a sector-applicability correction, NOT a "
        "judgement uplift - the uplift is argued separately in Stage 2.")

    # ------------------------------------------------------------- committee overlay (Stage 2)
    # Each notch is named and justified; the sheet computes the bridge from the A+ anchor to the
    # assigned AA+. Total +3.
    OVERLAY = {
        60: (0, "No further sector-metric adjustment. The one such adjustment on this case - "
                "the accrual-coverage cap - is handled at Inputs!C67/D67 above and is already "
                "reflected in the Stage 1 anchor. Counting it again here would double-count it."),
        61: (2, "+2 for franchise scale and diversification the financial composite cannot see. "
                "Reported revenue of Rs 5,131.43 crore understates the business by roughly 6.7 "
                "times: FY2026 Booking Value was Rs 34,171 crore (17,513 units, 27 msf) - the "
                "highest recorded by any Indian developer for the third consecutive year, "
                "corroborated across two independent annual reports - with collections of "
                "Rs 19,965 crore, +17%. The portfolio spans 10+ cities. The scorecard reads "
                "only the recognised-revenue line and therefore cannot see any of this."),
        62: (1, "+1 for Godrej Industries Group financial flexibility BEYOND the single notch "
                "already entered at Step 12b (Inputs!C42), stated here as the increment so "
                "support is not counted twice. Evidence: Rs 12,850 crore of equity raised "
                "FY2019-FY2025 through QIP/private placement, an access to capital ICRA's own "
                "published rationale describes as 'exceptional financial flexibility'. Held to "
                "+1 rather than more because it is demonstrated market access, not an explicit "
                "guarantee or a cash-support undertaking."),
        63: (0, "No peer-benchmarking adjustment. The resulting AA+ sits consistently with this "
                "agency's published grades for large listed Indian developers of comparable "
                "scale and leverage; no calibration gap was identified that the factors above "
                "do not already carry."),
        64: (0, "None."),
    }
    for row, (notches, why) in OVERLAY.items():
        ws.cell(row=row, column=3, value=notches)
        ws.cell(row=row, column=4, value=why)

    wb.save(DST)
    print("wrote %s" % DST)
    print("\nHistoricals fed in (consolidated, Rs crore, PAT REPORTED):")
    for y in HIST:
        print("  %-8s Rev %9s  OPBDIT %9s  PAT %9s  Debt %9s  NW %9s"
              % (y, "{:,.0f}".format(R[y]["Revenue"]),
                 "{:,.0f}".format(R[y]["OPBDIT (ex Other Income)"]),
                 "{:,.0f}".format(R[y]["PAT (total)"]),
                 "{:,.0f}".format(R[y]["Total debt"]),
                 "{:,.0f}".format(R[y]["Net worth"])))
    print("\nProjections copied from forward_model.py (base case):")
    for k, v in proj.items():
        print("  %-18s %s" % (k, v))
    print("\nQualitative scores and their justification:")
    for k in ("business", "management", "support", "liquidity"):
        v, why = QUALITATIVE[k]
        print("\n  %s = %s\n    %s" % (k, v, why))


if __name__ == "__main__":
    main()
