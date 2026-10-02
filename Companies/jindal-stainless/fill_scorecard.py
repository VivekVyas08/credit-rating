"""Step 12b - fill the native Rating_Scorecard template for Jindal Stainless Limited.

The template carries its own Inputs -> Ratios -> Scorecard -> Rating formula chain. This
script fills only the input cells and then reads the anchor grade back out of Rating!C12,
rather than building a parallel scorecard - the point of the exercise is an INDEPENDENT
quantitative cross-check on the judgement view, and a bespoke scorecard would not be
independent of the analyst who built it.

Historicals come from compute_ratios.py and projections are COPIED (never linked) from the
Step 12 model, each with its source noted, per CLAUDE.md Step 12(e).
"""

import shutil
import sys

from openpyxl import load_workbook

import compute_ratios as C

SRC = "../../Rating_Scorecard.xlsx"
DST = "Rating_Scorecard_Jindal Stainless Limited.xlsx"

HIST = ["FY2024", "FY2025", "FY2026"]          # template's three historical columns C, D, E


def net_block(y):
    n = C.NON_CURRENT_ASSETS
    return sum(n[k][C.IDX[y]] for k in ("Property, plant and equipment",
                                        "Capital work-in-progress",
                                        "Right-of-use assets", "Investment property"))


def capex(y):
    """Derived from the net-block roll-forward - the AR gives no single capex line."""
    prior = C.YEARS[C.IDX[y] - 1]
    return net_block(y) - net_block(prior) + C.PL["Depreciation"][C.IDX[y]]


# Dividend paid, consolidated cash-flow statement (financing activities).
DIVIDEND = {"FY2024": 285.73, "FY2025": 244.86, "FY2026": 244.86}

# Base case copied from Jindal Stainless Limited_5yr_Model.xlsx, evaluated 26-Aug-2026.
PROJ = {
    "revenue_growth": (0.0925, 0.0953),        # FY2027E, FY2028E
    "ebitda_margin": (0.128, 0.130),
    "capex": (4000.0, 3600.0),
    "net_new_debt": (355.0, -587.0),           # derived from the funding gap, not assumed
    "interest_rate": (0.078, 0.078),
    "tax_rate": (0.25, 0.25),
    "dividend_payout": (0.12, 0.12),
}

QUALITATIVE = {
    "business": (2, "India's largest stainless steel producer with ~2.57 MT of finished-goods "
                    "sales in FY2026 and group melt capacity reaching 4.2 MTPA, backward "
                    "integrated and now with an offshore melt platform in Indonesia. Scale, "
                    "integration and a dominant domestic share support a strong score; it is "
                    "not a 1 because stainless remains import-exposed and the product is "
                    "substantially commoditised."),
    "management": (2, "Unmodified audit opinion from joint statutory auditors in every year "
                      "reviewed, no qualification, Emphasis of Matter or going-concern flag "
                      "identified in the materiality sweep. Scored 2 rather than 1 given the "
                      "density of related-party dealing with other Jindal group entities "
                      "(Jindal United Steel, Jindal Coke), which is disclosed and appears "
                      "arm's length but is a governance factor a rating committee would test."),
    "support": (0, "No parent or group support notch. JSL is the listed parent of its own "
                   "group; support runs from JSL to its subsidiaries, not toward it."),
    "liquidity": ("Strong",
                  "Unencumbered liquid assets of Rs 2,931.60 crore at 31-Mar-2026 - cash and "
                  "equivalents 415.12, other bank balances 2,518.91 and current investments "
                  "1.12, LESS Rs 3.55 crore disclosed as under lien with banks. The "
                  "encumbrance is immaterial at 0.1% of the total, and it was checked rather "
                  "than assumed (failure mode 34: cash is not liquidity until shown "
                  "unencumbered). Against that, PBILDT of Rs 5,572.04 crore covers finance "
                  "costs 9.81 times and the funding-gap model shows debt peaking in FY2027 "
                  "and falling thereafter."),
    "outlook": ("Stable", "See the rationale."),
}


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                                     # noqa: BLE001
        pass

    shutil.copyfile(SRC, DST)
    wb = load_workbook(DST)
    ws = wb["Inputs"]
    R = C.compute()

    ws["C4"] = "Jindal Stainless Limited"
    ws["C5"] = "CareEdge"
    ws["C6"] = "Iron & Steel / Metals / Mining"
    ws["C7"] = "26-Aug-2026"

    def liquid(y):
        return C.liquid_assets(y)

    rows = {
        12: lambda y: C.PL["Revenue from ops"][C.IDX[y]],
        13: lambda y: R[y]["PBILDT"],
        14: lambda y: C.PL["Depreciation"][C.IDX[y]],
        15: lambda y: C.PL["Finance costs"][C.IDX[y]],
        16: lambda y: C.PL["PAT"][C.IDX[y]],
        17: lambda y: C.total_debt(y),
        18: liquid,
        19: lambda y: C.tangible_net_worth(y),
        20: lambda y: C.tol(y),
        21: lambda y: C.current_assets(y),
        22: lambda y: C.current_liabilities(y),
        23: capex,
        24: lambda y: DIVIDEND[y],
    }
    for r, fn in rows.items():
        for j, y in enumerate(HIST):
            ws.cell(row=r, column=3 + j, value=round(fn(y), 2))

    for r, key in [(29, "revenue_growth"), (30, "ebitda_margin"), (31, "capex"),
                   (32, "net_new_debt"), (33, "interest_rate"), (34, "tax_rate"),
                   (35, "dividend_payout")]:
        for j in range(2):
            ws.cell(row=r, column=3 + j, value=PROJ[key][j])

    ws["C39"] = None                                   # no industry override; sector base = 4
    ws["C40"] = QUALITATIVE["business"][0]
    ws["C41"] = QUALITATIVE["management"][0]
    ws["C42"] = QUALITATIVE["support"][0]
    ws["C43"] = QUALITATIVE["liquidity"][0]
    ws["C44"] = QUALITATIVE["outlook"][0]

    wb.save(DST)
    print("wrote %s" % DST)

    # Read the anchor back through the template's OWN formula chain.
    try:
        import _xlread
        cells = _xlread.load(DST)
        anchor = cells.get(("RATING", "C", 12))
        print("Rating!C12 anchor grade: %s" % anchor)
    except Exception as e:                                                # noqa: BLE001
        print("anchor read-back unavailable (%s); open the workbook to read Rating!C12" % e)

    print("\nQualitative scores and their justification:")
    for k in ("business", "management", "support", "liquidity"):
        v, why = QUALITATIVE[k]
        print("  %-11s %-9s %s" % (k, v, why))


if __name__ == "__main__":
    main()
