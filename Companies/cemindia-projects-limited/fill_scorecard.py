"""Step 12b - fill the native Rating_Scorecard template for Cemindia Projects Limited.

Historicals from compute_ratios.py; projections copied (never linked) from forward_model.py.
Stage 1 (anchor) only is filled here; Stage 2 (committee overlay) and Stage 3 (assigned rating)
are filled at Step 17c by the committee, then annotate_scorecard.py verifies the match.

No hard-cap disapplication is needed: accrual interest coverage is 4.99x (FY2026), so the
coverage cap does not fire, and net worth is strongly positive. Business risk is a STRENGTH
here (large scale, diversification, order book, Adani parentage), so the business-risk gate
(failure mode 53) should not bind - it is checked, not assumed.
"""

import os
import shutil
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import compute_ratios as C                                              # noqa: E402
import forward_model as F                                               # noqa: E402

SRC = "../../Rating_Scorecard.xlsx"
DST = "Rating_Scorecard_Cemindia Projects Limited.xlsx"
HIST = ["FY2024", "FY2025", "FY2026"]

# Cash capex (purchase of PPE/intangibles), consolidated cash-flow statement. FY2024 not cleanly
# available on the old ITD basis in crore; left blank rather than converted imprecisely.
CAPEX_CASHFLOW = {"FY2024": None, "FY2025": 228.05, "FY2026": 263.80}
DIVIDEND = {"FY2024": None, "FY2025": 29.16, "FY2026": 34.30}   # cash-flow financing block

QUALITATIVE = {
    "business": (2, "Large, diversified EPC/construction contractor with FY2026 consolidated "
                    "revenue of Rs 10,061 crore (+8.4% YoY, five-year CAGR 30%) and a strong, "
                    "well-diversified order book of Rs 31,306.8 crore as on June 30, 2026 (3.1x "
                    "FY2026 OI, up ~71% from Rs 18,300 crore at March 2025) - an over-four-decade "
                    "engineering franchise (incorporated June 1978; formerly ITD Cementation) "
                    "with recognised leadership in marine/maritime and tunnelling/underground "
                    "metro works. The book is diversified across marine (23.2%), urban "
                    "infrastructure/MRTS (16.5%), data centres (13.3%), water/sewage (10.9%), "
                    "buildings (8.8%), transportation (8.4%) and industrial (8.1%). Now "
                    "67.46%-held by RenewExim DMCC (Adani group): CPL is the group's key EPC arm, "
                    "with ~65% of ~Rs 25,000 crore of new orders over the 15 months to June 2026 "
                    "from the group, plus favourable BG terms and bank rates. Scored 2, not 3, "
                    "because scale, a 3.1x diversified book and demonstrated Adani synergies are "
                    "genuine strengths; not 1, because EPC is competitively intense, the book is "
                    "in early execution stages, and margins are structurally in the 10-11% band."),
    "management": (3, "Statutory auditors expressed an UNMODIFIED opinion on both the "
                      "standalone and consolidated FY2026 financial statements (Price "
                      "Waterhouse Chartered Accountants LLP, 29-Apr-2026). Adani-group "
                      "governance and a stable senior team post the FY2025 change of control. "
                      "Held at 3 rather than 2 because of a CARO clause 3(ii)(b) observation - "
                      "quarterly working-capital returns filed with banks not in agreement "
                      "with the unaudited books, and the March-2026-quarter returns not filed "
                      "- a process/governance matter (not a financial-statement qualification "
                      "or going-concern flag), and because the change of control is recent."),
    "support": (0, "NOT APPLICABLE as a support notch. RenewExim DMCC (Adani group) holds "
                   "67.46%, and the Adani association is a material credit positive - but it "
                   "operates through BUSINESS channels (sizeable order inflows, favourable BG "
                   "terms and bank rates, access to the group's project pipeline), not through "
                   "an explicit guarantee or cash-support undertaking. ICRA's own action treats "
                   "Parent/Group support as 'Not Applicable' and captures the benefit inside "
                   "business risk and financial flexibility. Entering zero here keeps the anchor "
                   "honestly standalone; the group-driven uplift is applied as a named Stage-2 "
                   "business-strength overlay (see the committee overlay), where it is visible, "
                   "rather than smuggled in as a support notch the methodology does not support."),
    "liquidity": ("Adequate",
                  "Free cash and bank balances of around Rs 380 crore plus a buffer of around "
                  "Rs 209 crore in fund-based working-capital limits as on June 30, 2026, "
                  "against modest debt repayment obligations of Rs 61.1 crore in FY2027 and "
                  "Rs 49.4 crore in FY2028, and healthy net cash accruals of around Rs 800-1,000 "
                  "crore per annum. Annual capex is around Rs 350-400 crore with no material "
                  "debt-funding plans, and shareholder approval exists for a QIP of up to "
                  "Rs 5,000 crore that would further limit reliance on external debt. Adequate "
                  "rather than Strong because the operations are working-capital-intensive "
                  "(inventory mainly WIP, NWC/OI rising to 4.7% in FY2026) and are funded partly "
                  "through mobilisation advances of ~Rs 1,015 crore and extended supplier credit "
                  "whose continuation matters; Adequate rather than Stretched because accruals "
                  "cover the scheduled repayments many times over and there is no debt cliff."),
    "outlook": ("Stable", "See the rating view in the workpaper."),
}


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                                     # noqa: BLE001
        pass

    shutil.copyfile(SRC, DST)
    wb = load_workbook(DST)
    ws = wb["Inputs"]
    R = C.compute()
    I = {y: i for i, y in enumerate(C.YEARS)}
    base = F.project(False)

    ws["C4"] = "Cemindia Projects Limited"
    ws["C5"] = "ICRA"
    ws["C6"] = "Real Estate / Construction"    # nearest of the 22 scorecard sectors to EPC
    ws["C7"] = "19-Sep-2026"

    rows = {
        12: lambda y: R[y]["Operating income"],
        13: lambda y: R[y]["OPBDIT"],
        14: lambda y: C.CON["depreciation"][I[y]],
        15: lambda y: C.CON["finance_cost"][I[y]],
        16: lambda y: R[y]["PAT"],
        17: lambda y: R[y]["Total debt"],
        18: lambda y: R[y]["Unencumbered cash"],
        19: lambda y: R[y]["Tangible net worth"],
        20: lambda y: C.CON["total_assets"][I[y]] - C.CON["equity"][I[y]],
        21: lambda y: C.CON["curr_assets"][I[y]],
        22: lambda y: C.CON["curr_liab"][I[y]],
        23: lambda y: CAPEX_CASHFLOW[y],
        24: lambda y: DIVIDEND[y],
    }
    for r, fn in rows.items():
        for j, y in enumerate(HIST):
            v = fn(y)
            if v is not None:
                ws.cell(row=r, column=3 + j, value=round(v, 2))

    fy26_rev = R["FY2026"]["Operating income"]
    proj = {
        "revenue_growth": (round(base[0]["Revenue"] / fy26_rev - 1, 4),
                           round(base[1]["Revenue"] / base[0]["Revenue"] - 1, 4)),
        "ebitda_margin": (base[0]["OPBDIT margin %"] / 100.0, base[1]["OPBDIT margin %"] / 100.0),
        "capex": (base[0]["Capex"], base[1]["Capex"]),
        "net_new_debt": (base[0]["New debt raised"], base[1]["New debt raised"]),
        "interest_rate": (0.10, 0.10),
        "tax_rate": (0.27, 0.27),
        "dividend_payout": (0.12, 0.12),
    }
    for r, key in [(29, "revenue_growth"), (30, "ebitda_margin"), (31, "capex"),
                   (32, "net_new_debt"), (33, "interest_rate"), (34, "tax_rate"),
                   (35, "dividend_payout")]:
        for j in range(2):
            ws.cell(row=r, column=3 + j, value=proj[key][j])

    ws["C39"] = None
    ws["C40"] = QUALITATIVE["business"][0]
    ws["C41"] = QUALITATIVE["management"][0]
    ws["C42"] = QUALITATIVE["support"][0]
    ws["C43"] = QUALITATIVE["liquidity"][0]
    ws["C44"] = QUALITATIVE["outlook"][0]

    wb.save(DST)
    print("wrote %s" % DST)
    for y in HIST:
        print("  %-8s OI %9s  OPBDIT %8s  PAT %8s  Debt %8s  TNW %8s" % (
            y, "{:,.0f}".format(R[y]["Operating income"]), "{:,.0f}".format(R[y]["OPBDIT"]),
            "{:,.0f}".format(R[y]["PAT"]), "{:,.0f}".format(R[y]["Total debt"]),
            "{:,.0f}".format(R[y]["Tangible net worth"])))
    print("\nQualitative:")
    for k in ("business", "management", "support", "liquidity"):
        print("  %-11s %s" % (k, QUALITATIVE[k][0]))


if __name__ == "__main__":
    main()
