"""Step 12b - fill the native Rating_Scorecard template for DOMS Industries Limited (CRISIL).

Historicals are read from compute_ratios.py (never retyped). Projections are COPIED, not
linked, from forward_model.py, which derives the debt path from the capex funding gap.

ONE SECTOR DECISION IS LOAD-BEARING AND IS STATED HERE RATHER THAN BURIED. The template offers
22 sectors and none of them is "stationery". The closest structural match is FMCG / Consumer
staples: DOMS is a branded consumer-products manufacturer selling through roughly 145,000
retail outlets, with the demand stability that implies. But the staples base industry score of
2 is too generous for this business, and the reason is on the face of the FY2027 accounts -
raw-material inflation took the quarterly OPBDITA margin from 17.37% to 12.08% in a single
year. A staples producer with pricing power does not do that. The INDUSTRY OVERRIDE cell is
therefore set to 3 rather than accepting the sector base, and this note is the justification
the override requires.
"""

import os
import shutil
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import compute_ratios as C                                              # noqa: E402
import forward_model as F                                               # noqa: E402

SRC = "../../Rating_Scorecard.xlsx"
DST = "Rating_Scorecard_DOMS Industries Limited.xlsx"
HIST = ["FY2024", "FY2025", "FY2026"]
LAKH = 100.0

# CONSOLIDATED cash-flow statements. FY2024 and FY2025 come from the FY2024-25 annual report's
# CONSOLIDATED statement (L10850-10851, L10881); FY2026 from the FY2025-26 annual report
# (L11489-11490, L11503). The consolidated statement is used deliberately: the same annual
# report's STANDALONE cash flow shows FY2025 PPE purchases of 18,313.12 lakh against the
# consolidated 21,002.37 lakh, and taking the standalone figure into a consolidated model
# would understate the programme by Rs 26.9 crore.
CAPEX_CASHFLOW = {
    "FY2024": (15280.56 + 65.13) / LAKH,
    "FY2025": (21002.37 + 325.00) / LAKH,
    "FY2026": (28078.30 + 1205.24) / LAKH,
}
DIVIDEND = {
    "FY2024": 931.30 / LAKH,
    "FY2025": 1517.06 / LAKH,
    "FY2026": 1911.62 / LAKH,
}

QUALITATIVE = {
    "industry_override": (3, "See the module docstring. FMCG / Consumer staples is the closest "
                             "of the 22 template sectors, but its base score of 2 assumes "
                             "pricing power this business has not demonstrated: the Q1 FY2027 "
                             "OPBDITA margin fell 529 basis points year on year on "
                             "raw-material inflation, and the company states it chose "
                             "volume-led growth and market share over near-term margin. "
                             "Overridden to 3."),
    "business": (2, "The clear number-two in Indian stationery with a portfolio spanning "
                    "wooden pencils, writing instruments, mechanical pencils, paper "
                    "stationery, art materials and adhesives, distributed through roughly "
                    "145,000 retail outlets, plus a Hygiene segment (Rs 203.10 crore FY2026) "
                    "that grew 81% and gives genuine category diversification. Exports are "
                    "13-15% of sales and the Reynolds brand has been added. Scored 2 rather "
                    "than 3 because the franchise is deep and the growth is volume-led rather "
                    "than price-led: revenue compounded 24.4% then 21.6% over two years, and "
                    "the SKU count rose by about 400 in a single year. Not a 1, because the "
                    "category is discretionary at the margin and the company competes with "
                    "one much larger domestic incumbent."),
    "management": (2, "Unmodified audit opinion. The entire Rs 332.72 crore of IPO net "
                      "proceeds is disclosed as fully utilised in accordance with the objects "
                      "in the offer document, with a line-by-line reconciliation - which is "
                      "the specific test of whether stated intentions were followed. The "
                      "capital programme is being executed as described (Umbergaon greenfield "
                      "phased, 300,000 sq ft operational in FY2027) and the company disclosed "
                      "the margin compression and its cause plainly rather than presenting it "
                      "as a one-off. Scored 2, not 1, because the shareholders' agreement "
                      "grants special rights to the promoter group that required a separate "
                      "member approval (86.28% in favour), which is a governance feature "
                      "worth naming rather than a defect."),
    "support": (0, "ZERO NOTCHES, and this is a deliberate reading rather than an omission. "
                   "F.I.L.A. - Fabbrica Italiana Lapis ed Affini S.p.A. is a FOREIGN "
                   "CO-PROMOTER holding 26.01% within a promoter group of 70.39%; the Indian "
                   "promoters hold 44.38%. FILA is also a commercial counterparty - exports to "
                   "the FILA Group are a named revenue channel. A 26% strategic shareholder "
                   "that is simultaneously a customer is not a parent providing implicit "
                   "support in the sense the support criteria mean, and notching for it would "
                   "be an invention. The company does not need it: borrowings repay in full by "
                   "FY2029 in the base case."),
    "liquidity": ("Strong",
                  "The descriptor is the CONCLUSION of the arithmetic, not an adjective. "
                  "Sources over twelve months: FY2026 net cash from operations of Rs 254.34 "
                  "crore (consolidated cash flow), cash and equivalents of Rs 53.49 crore, and "
                  "other bank balances of Rs 8.30 crore. Claims: current borrowings of Rs "
                  "51.55 crore, current lease liabilities of Rs 15.45 crore, and the FY2027 "
                  "capex the forward model derives at Rs 350 crore. Operating cash flow alone "
                  "covers the current debt maturities 3.8 times over. Strong rather than "
                  "Superior because the FY2027 capex exceeds operating cash flow at the "
                  "modelled 14.0% margin, so the year is funded partly by borrowing - the "
                  "model derives Rs 79 crore of new debt - and because the margin recovery "
                  "that closes the gap from FY2028 is an assumption, not yet an outcome. "
                  "Strong rather than Adequate because total debt including leases is Rs "
                  "141.26 crore against a tangible net worth of Rs 1,219.73 crore, gearing is "
                  "0.12x, and interest cover is 37x."),
    "outlook": ("Stable", "See the rating view in the workpaper."),
}


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                                     # noqa: BLE001
        pass

    shutil.copyfile(SRC, DST)
    wb = load_workbook(DST)
    ws = wb["Inputs"]
    D = C.DATA
    I = {y: i for i, y in enumerate(C.YEARS)}
    base = F.project(False)

    ws["C4"] = "DOMS Industries Limited"
    ws["C5"] = "CRISIL"
    ws["C6"] = "FMCG / Consumer staples"
    ws["C7"] = "12-Aug-2026"

    def ebitda(y):
        i = I[y]
        return (D["pbt"][i] + D["finance_cost"][i] + D["depreciation"][i]
                - D["other_income"][i] - D["assoc_share"][i])

    rows = {
        12: lambda y: D["revenue"][I[y]],
        13: ebitda,
        14: lambda y: D["depreciation"][I[y]],
        15: lambda y: D["finance_cost"][I[y]],
        # PAT to OWNERS, consistent with tangible net worth being taken net of the
        # non-controlling interest. The NCI is material here (Rs 82.37 crore of equity,
        # Rs 9.38 crore of FY2026 profit), so the two must be kept on the same basis.
        16: lambda y: D["pat_owners"][I[y]],
        17: lambda y: (D["borrow_nc"][I[y]] + D["borrow_c"][I[y]]
                       + D["lease_nc"][I[y]] + D["lease_c"][I[y]]),
        18: lambda y: D["cash"][I[y]] + D["bank_other"][I[y]],
        19: lambda y: D["equity_total"][I[y]] - D["nci_equity"][I[y]],
        20: lambda y: D["total_assets"][I[y]] - D["equity_total"][I[y]],
        21: lambda y: D["curr_assets"][I[y]],
        22: lambda y: D["curr_liab"][I[y]],
        23: lambda y: CAPEX_CASHFLOW[y],
        24: lambda y: DIVIDEND[y],
    }
    for r, fn in rows.items():
        for j, y in enumerate(HIST):
            ws.cell(row=r, column=3 + j, value=round(fn(y), 2))

    fy26 = D["revenue"][2]
    proj = {
        "revenue_growth": (round(base[0]["Revenue"] / fy26 - 1, 4),
                           round(base[1]["Revenue"] / base[0]["Revenue"] - 1, 4)),
        "ebitda_margin": (base[0]["OPBDITA margin %"] / 100.0,
                          base[1]["OPBDITA margin %"] / 100.0),
        "capex": (base[0]["Capex"], base[1]["Capex"]),
        "net_new_debt": (base[0]["New debt raised"], base[1]["New debt raised"]),
        "interest_rate": (0.065, 0.065),
        "tax_rate": (0.257, 0.257),
        "dividend_payout": (0.09, 0.09),
    }
    for r, key in [(29, "revenue_growth"), (30, "ebitda_margin"), (31, "capex"),
                   (32, "net_new_debt"), (33, "interest_rate"), (34, "tax_rate"),
                   (35, "dividend_payout")]:
        for j in range(2):
            ws.cell(row=r, column=3 + j, value=proj[key][j])

    ws["C39"] = QUALITATIVE["industry_override"][0]
    ws["C40"] = QUALITATIVE["business"][0]
    ws["C41"] = QUALITATIVE["management"][0]
    ws["C42"] = QUALITATIVE["support"][0]
    ws["C43"] = QUALITATIVE["liquidity"][0]
    ws["C44"] = QUALITATIVE["outlook"][0]

    wb.save(DST)
    print("wrote %s" % DST)
    print("\nHistoricals fed in (consolidated, Rs crore; PAT to owners; debt incl. leases):")
    for y in HIST:
        print("  %-8s OI %9s  OPBDITA %8s  PAT %8s  Debt %8s  TNW %9s  Capex %7s"
              % (y, "{:,.1f}".format(rows[12](y)), "{:,.1f}".format(ebitda(y)),
                 "{:,.1f}".format(rows[16](y)), "{:,.1f}".format(rows[17](y)),
                 "{:,.1f}".format(rows[19](y)), "{:,.1f}".format(rows[23](y))))
    print("\nProjections copied from forward_model.py (base case, FY2027 / FY2028):")
    for k, v in proj.items():
        print("  %-18s %s" % (k, v))
    print("\nQualitative scores and their justification:")
    for k in ("industry_override", "business", "management", "support", "liquidity"):
        v, why = QUALITATIVE[k]
        print("\n  %s = %s\n    %s" % (k, v, why))


if __name__ == "__main__":
    main()
