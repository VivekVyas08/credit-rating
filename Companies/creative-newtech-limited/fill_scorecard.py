"""Step 12b - fill the native Rating_Scorecard template for Creative Newtech Limited (CRISIL).

Historicals are read from compute_ratios.py (never retyped). Projections are COPIED, not
linked, from forward_model.py, which derives the debt path from the working-capital funding
gap rather than assuming it.

TWO THINGS ARE STATED HERE RATHER THAN BURIED.

1. FY2024 IS DELIBERATELY LEFT PARTLY BLANK. The engine holds no FY2024 figure for
   depreciation, other income, PAT or cash - they were never extracted, and the FY2024 column
   of the audited FY2026 filing does not carry them. They are left EMPTY rather than filled
   with a plausible number. This costs nothing analytically: the scorecard's financial score
   reads only the latest actual year (Ratios column E) and the FY+1 projection (column F), so
   FY2024 does not reach the composite at any point.

2. THE SOURCE CASH FLOW IS INTERNALLY INCONSISTENT AND IS NOT SILENTLY RELIED ON. In the
   audited FY2026 consolidated cash-flow statement, depreciation is stated as 30.43 lakh
   against a P&L charge of 152.10 lakh, and finance cost as 1,746.91 lakh against a P&L charge
   of 2,066.43 lakh. The FY2025 depreciation reconciles exactly (116.07 lakh both places), so
   this is an FY2026 presentation problem, not a definitional one. Only the lines that are
   unambiguous - net operating cash flow, purchase of PPE, dividend paid - are used, and the
   discrepancy is recorded here rather than left for a reader to trip over.
"""

import os
import shutil
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import compute_ratios as C                                              # noqa: E402
import forward_model as F                                               # noqa: E402

SRC = "../../Rating_Scorecard.xlsx"
DST = "Rating_Scorecard_Creative Newtech Limited.xlsx"
HIST = ["FY2024", "FY2025", "FY2026"]
LAKH = 100.0

# Consolidated cash-flow statement, audited FY2026 results filing (OutcomeBM14052026-1.md).
# FY2024 was not extracted - left None, not estimated.
CAPEX_CASHFLOW = {"FY2024": None, "FY2025": 281.66 / LAKH, "FY2026": 330.04 / LAKH}
DIVIDEND = {"FY2024": None, "FY2025": 71.21 / LAKH, "FY2026": 75.08 / LAKH}

QUALITATIVE = {
    "business": (4, "A value-added distributor of consumer technology and gadgets with four "
                    "reported segments, transitioning into brand licensing and contract "
                    "manufacturing (Honeywell). Revenue grew 52.1% to Rs 2,699.77 crore in "
                    "FY2026. Scored 4 rather than 3 because the business is intermediary by "
                    "nature - purchases of stock-in-trade are 94.5% of revenue and gross "
                    "margin is 7.91%, so the company earns a thin spread on other people's "
                    "brands and has limited pricing power. Not a 5, because the segment base "
                    "is genuinely diversified, the brand-licensing arm gives a route to a "
                    "structurally better margin, and scale has grown 2.6 times in two years."),
    "management": (4, "The FY2026 audit opinion is unmodified. Scored 4 rather than 3 on two "
                      "specific grounds, both evidenced: receivable days went from 49 to 76 in "
                      "a single year while revenue grew 52.1%, which is a working-capital "
                      "discipline question rather than a growth achievement; and the FY2026 "
                      "consolidated cash-flow statement carries a depreciation figure "
                      "(30.43 lakh) that does not reconcile to its own P&L (152.10 lakh). "
                      "Not a 5 - no qualification, no going-concern matter, no auditor "
                      "observation, and the dividend has been maintained."),
    "support": (0, "No parent and no group support. The company is promoter-managed and "
                   "stands on its own balance sheet, so a support notch would be an invention. "
                   "Entered as zero, which keeps the anchor standalone."),
    "liquidity": ("Stretched",
                  "This descriptor is the CONCLUSION of the arithmetic below, not an "
                  "adjective. Claims over twelve months: the entire Rs 324.25 crore of debt is "
                  "current and repayable or renewable within the year, plus the FY2027 "
                  "working-capital build the forward model derives at Rs 90 crore. Sources: "
                  "cash and cash equivalents of Rs 5.36 crore; other bank balances of Rs 16.82 "
                  "crore, which include margin money placed against facilities (the FY2026 "
                  "movement alone was Rs 5.47 crore) and are therefore NOT shown to be "
                  "unencumbered; and operating cash flow which in FY2026 was NEGATIVE Rs "
                  "273.10 crore. The company funded that burn with Rs 254.72 crore of fresh "
                  "short-term borrowing and by running cash down 87.8%. Stretched rather than "
                  "Poor because the borrowing is working-capital limits against a receivable "
                  "and inventory book of Rs 738.56 crore that comfortably covers them, the "
                  "current ratio is 1.68x, and the business is profitable - Rs 61.75 crore to "
                  "owners in FY2026. Stretched rather than Adequate because liquidity here "
                  "depends entirely on lenders continuing to roll limits that grew 4.5 times "
                  "in one year, and on collecting a receivable book that grew 137.7%."),
    "outlook": ("Stable", "See the rating view in the workpaper."),
}


def val(fn, y):
    """Return a rounded figure, or None where the input was never extracted."""
    try:
        v = fn(y)
    except (TypeError, KeyError):
        return None
    return None if v is None else round(v, 2)


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                                     # noqa: BLE001
        pass

    shutil.copyfile(SRC, DST)
    wb = load_workbook(DST)
    ws = wb["Inputs"]
    D = C.DATA
    I = {y: i for i, y in enumerate(C.YEARS)}
    base = F.project(False)

    ws["C4"] = "Creative Newtech Limited"
    ws["C5"] = "CRISIL"
    ws["C6"] = "Trading / Distribution"
    ws["C7"] = "12-Aug-2026"

    def oi(y):
        return D["revenue"][I[y]] + D["export_inc"][I[y]]

    def ebitda(y):
        i = I[y]
        if D["depreciation"][i] is None or D["other_income"][i] is None:
            return None
        return (D["pbt"][i] + D["finance_cost"][i] + D["depreciation"][i]
                - D["other_income"][i])

    def cash(y):
        i = I[y]
        if D["cash"][i] is None:
            return None
        return D["cash"][i] + D["bank_other"][i]

    rows = {
        12: oi,
        13: ebitda,
        14: lambda y: D["depreciation"][I[y]],
        15: lambda y: D["finance_cost"][I[y]],
        # PAT to OWNERS, consistent with tangible net worth being taken net of the
        # non-controlling interest. Using total PAT against an owners-only net worth would
        # flatter every return and accrual ratio on the sheet.
        16: lambda y: D["pat_owners"][I[y]],
        17: lambda y: D["borrow_nc"][I[y]] + D["borrow_c"][I[y]],
        18: cash,
        19: lambda y: D["equity_total"][I[y]] - D["nci_equity"][I[y]],
        20: lambda y: D["total_assets"][I[y]] - D["equity_total"][I[y]],
        21: lambda y: D["curr_assets"][I[y]],
        22: lambda y: D["curr_liab"][I[y]],
        23: lambda y: CAPEX_CASHFLOW[y],
        24: lambda y: DIVIDEND[y],
    }
    for r, fn in rows.items():
        for j, y in enumerate(HIST):
            ws.cell(row=r, column=3 + j, value=val(fn, y))

    fy26_oi = oi("FY2026")
    proj = {
        "revenue_growth": (round(base[0]["Revenue"] / fy26_oi - 1, 4),
                           round(base[1]["Revenue"] / base[0]["Revenue"] - 1, 4)),
        "ebitda_margin": (base[0]["OPBDITA margin %"] / 100.0,
                          base[1]["OPBDITA margin %"] / 100.0),
        "capex": (base[0]["Capex"], base[1]["Capex"]),
        "net_new_debt": (base[0]["New debt raised"], base[1]["New debt raised"]),
        "interest_rate": (0.105, 0.105),
        "tax_rate": (0.22, 0.22),
        # Rs 0.75 crore against projected PAT of Rs 63.3 crore. The Board has recommended
        # Re 0.50 per Rs 10 share for FY2026; the payout is nominal and is held flat in rupees
        # rather than as a percentage, so the ratio falls as earnings grow.
        "dividend_payout": (round(F.BASE["dividend"] / base[0]["PAT"], 4),
                            round(F.BASE["dividend"] / base[1]["PAT"], 4)),
    }
    for r, key in [(29, "revenue_growth"), (30, "ebitda_margin"), (31, "capex"),
                   (32, "net_new_debt"), (33, "interest_rate"), (34, "tax_rate"),
                   (35, "dividend_payout")]:
        for j in range(2):
            ws.cell(row=r, column=3 + j, value=proj[key][j])

    ws["C39"] = None                     # no override; Trading / Distribution sector base = 4
    ws["C40"] = QUALITATIVE["business"][0]
    ws["C41"] = QUALITATIVE["management"][0]
    ws["C42"] = QUALITATIVE["support"][0]
    ws["C43"] = QUALITATIVE["liquidity"][0]
    ws["C44"] = QUALITATIVE["outlook"][0]

    wb.save(DST)
    print("wrote %s" % DST)
    print("\nHistoricals fed in (consolidated, Rs crore; PAT to owners):")
    for y in HIST:
        def s(fn):
            v = val(fn, y)
            return "n/a" if v is None else "{:,.1f}".format(v)
        print("  %-8s OI %9s  EBITDA %8s  PAT %8s  Debt %8s  TNW %8s  Cash %7s"
              % (y, s(oi), s(ebitda), s(rows[16]), s(rows[17]), s(rows[19]), s(cash)))
    print("\nProjections copied from forward_model.py (base case, FY2027 / FY2028):")
    for k, v in proj.items():
        print("  %-18s %s" % (k, v))
    print("\nQualitative scores and their justification:")
    for k in ("business", "management", "support", "liquidity"):
        v, why = QUALITATIVE[k]
        print("\n  %s = %s\n    %s" % (k, v, why))


if __name__ == "__main__":
    main()
