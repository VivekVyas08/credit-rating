# -*- coding: utf-8 -*-
"""Generate Equitas_SFB_Financial_Spreads.xlsx FROM compute_ratios.py (R6 - the script is the source
of truth). Tabs: Extraction / RatioSuite / TieOut / ExternalTieOut / Notes. Values are written
literally and render-verified by a data_only read-back."""
import importlib.util
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("cr", os.path.join(HERE, "compute_ratios.py"))
cr = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cr)

HDR, TITLE = Font(bold=True), Font(bold=True, size=12)
FILL = PatternFill("solid", fgColor="DDEBF7")
OK, BAD = PatternFill("solid", fgColor="E2F3E6"), PatternFill("solid", fgColor="FDECE0")
WRAP = Alignment(wrap_text=True, vertical="top")
YEARS = cr.YEARS

SOURCES = {
    "BS_LIAB": ("Balance sheet - capital and liabilities, Rs '000 (FY2025-26 AR; FY2024-25 AR; FY2023-24 AR)", 1e-4),
    "BS_ASSET": ("Balance sheet - assets, Rs '000 (same reports)", 1e-4),
    "BS_PRINTED": ("Balance sheet printed total, Rs '000", 1e-4),
    "DEPOSITS": ("Schedule 3 - deposits, Rs '000 (FY2023 split only in MD&A)", 1e-4),
    "PL": ("Profit and loss account, Rs '000", 1e-4),
    "PL_PRINTED": ("Profit and loss printed totals, Rs '000", 1e-4),
    "INT_ON_ADVANCES": ("Schedule 13 - interest on advances, Rs '000", 1e-4),
    "PROV": ("Note 18.14(f) 'Provisions and Contingencies' (Rs crore)", 1),
    "NPA": ("Asset-quality note - classification of advances and provisions (Rs crore)", 1),
    "NPA_RATIOS_PRINTED": ("Asset-quality note - NPA ratios as printed (%)", 1),
    "SR": ("Investment notes - security receipts and provisions (Rs crore)", 1),
    "RESTRUCTURED_STD": ("Resolution-framework (COVID) table - restructured standard (Rs crore)", 1),
    "CAP": ("Note 1a Composition of Regulatory Capital, Basel II (Rs crore)", 1),
    "CAP_PRINTED": ("Note 1a ratios as printed (%)", 1),
    "TNW_ADJ": ("Schedule 11 DTA (Rs '000) and software note (Rs crore)", 1),
    "NETWORTH_PRINTED": ("Q4 FY2026 results Reg 52(4) net worth; FY2023-24 AR highlights", 1),
    "LIQ": ("LCR / NSFR / concentration notes (%)", 1),
    "BANK_RATIOS": ("Business ratios note - bank's own definitions (%)", 1),
    "MDA": ("FY2025-26 AR MD&A profit table; FY2023-24 AR text (Rs crore)", 1),
    "OPS": ("MD&A product tables / highlights / BRSR (Rs crore; counts)", 1),
    "INT_EARNED_SPLIT": ("Schedule 13 split, Rs '000", 1e-4),
    "LCR_COMPONENTS": ("LCR note, quarter ended March (Rs crore)", 1),
    "ROA_Q4_RESULTS": ("Q4 FY2026 results (%)", 1),
}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Extraction"
ws["A1"] = "Equitas Small Finance Bank Limited - extracted line items"
ws["A1"].font = TITLE
ws["A2"] = ("Literal copy of compute_ratios.py, converted to Rs crore where the source is Rs '000 (one "
            "conversion, in code). Blank = not disclosed in that year's documents / not extracted (Notes).")
ws["A2"].alignment = WRAP
ws.row_dimensions[2].height = 30
ws.append([])
ws.append(["Table", "Line item"] + YEARS + ["Unit as shown", "Source"])
for c in ws[4]:
    c.font, c.fill = HDR, FILL
for name, (src, mult) in SOURCES.items():
    for key, series in getattr(cr, name).items():
        m = mult if "Rs '000" not in key else 1e-4
        if name == "TNW_ADJ" and "Rs crore" in key:
            m = 1
        ws.append([name, key] + [None if x is None else round(float(x) * m, 4) for x in series]
                  + ["Rs crore" if m != 1 or "Rs '000" in src else "as printed", src])
ws.append([])
ws.append(["Scalar / partial-year inputs"])
ws.cell(ws.max_row, 1).font = HDR
for name in ("DEPOSITS_FY23_CR", "OPENING_GROSS_ADV_FY23", "TOP20_DEPOSITS_FY26", "PRODUCT_MIX", "MFI_SHARE_PRINTED_FY25",
             "SR_MOVEMENT", "RESTRUCTURED_MOVEMENT", "SOFTWARE_OPENING", "DTA_FY24_IN_FY25_AR", "QTR_PAT"):
    ws.append([name, str(getattr(cr, name))])
for col, w in zip("ABCDEFGH", [22, 44, 13, 13, 13, 13, 12, 70]):
    ws.column_dimensions[col].width = w

R = cr.compute()
ws2 = wb.create_sheet("RatioSuite")
ws2["A1"] = "CareEdge financial-sector ratio suite (February 2025 definitions) - standalone"
ws2["A1"].font = TITLE
ws2.append([])
ws2.append(["Metric"] + YEARS)
for c in ws2[3]:
    c.font, c.fill = HDR, FILL
keys = []
for y in YEARS:
    keys += [x for x in R[y] if x not in keys]
for key in keys:
    ws2.append([key] + [None if R[y].get(key) is None else round(float(R[y][key]), 2) for y in YEARS])
ws2.column_dimensions["A"].width = 56
for col in "BCDE":
    ws2.column_dimensions[col].width = 13

T = cr.tieouts(R)
ws3 = wb.create_sheet("TieOut")
ws3["A1"] = "Internal tie-outs - every check compares a computation with a figure the bank prints"
ws3["A1"].font = TITLE
ws3.append([])
ws3.append(["Check", "Computed", "Disclosed", "Difference", "Tolerance", "Result", "Source"])
for c in ws3[3]:
    c.font, c.fill = HDR, FILL
npass = 0
for d, comp, disc, tol, src in T:
    ok = abs(comp - disc) <= tol
    npass += ok
    ws3.append([d, round(comp, 4), round(disc, 4), round(comp - disc, 4), tol, "PASS" if ok else "FAIL", src])
    ws3.cell(row=ws3.max_row, column=6).fill = OK if ok else BAD
ws3.append([])
ws3.append(["TOTAL", None, None, None, None, "%d/%d PASS" % (npass, len(T))])
for col, w in zip("ABCDEFG", [74, 13, 13, 11, 10, 9, 56]):
    ws3.column_dimensions[col].width = w

ws4 = wb.create_sheet("ExternalTieOut")
ws4["A1"] = "External tie-out vs screener.in (cross-check ONLY - R1; read 10-Oct-2026)"
ws4["A1"].font = TITLE
ws4.append([])
ws4.append(["Item", "Year", "screener.in", "Engine", "Difference", "Result"])
for c in ws4[3]:
    c.font, c.fill = HDR, FILL
SCR = []
for y, rev, intx, pbt, np_, cap, res, dep, bor, ta in (
        ("FY2023", 4162, 1617, 769, 574, 1111, 4047, 25348, 2974, 34958),
        ("FY2024", 5486, 2407, 1063, 799, 1135, 4834, 36129, 1788, 45304),
        ("FY2025", 6312, 3060, 199, 147, 1140, 4933, 43102, 2137, 52836),
        ("FY2026", 6794, 3403, 128, 103, 1141, 4984, 46533, 5773, 60610)):
    SCR += [("Revenue (interest earned)", y, rev, R[y]["Interest earned"]),
            ("Interest expended", y, intx, cr.kv(cr.PL, "Interest expended", y)),
            ("Profit before tax", y, pbt, R[y]["PBT"]), ("Net profit", y, np_, R[y]["PAT"]),
            ("Equity capital + reserves", y, cap + res, R[y]["Net worth"]),
            ("Deposits", y, dep, R[y]["Deposits"]), ("Borrowings", y, bor, R[y]["Borrowings"]),
            ("Total assets", y, ta, R[y]["Total assets"])]
for item, y, s, e in SCR:
    diff = e - s
    ok = abs(diff) < 1.0
    ws4.append([item, y, s, round(e, 2), round(diff, 2), "PASS (rounding)" if ok else "CHECK"])
    ws4.cell(row=ws4.max_row, column=6).fill = OK if ok else BAD
for col, w in zip("ABCDEF", [30, 9, 12, 12, 11, 18]):
    ws4.column_dimensions[col].width = w

ws5 = wb.create_sheet("Notes")
ws5["A1"] = "Basis of preparation and disclosed differences"
ws5["A1"].font = TITLE
ws5.append([])
for title, text in cr.DISCLOSED_DIFFERENCES:
    ws5.append([title])
    ws5.cell(row=ws5.max_row, column=1).font = HDR
    ws5.append([text])
    ws5.cell(row=ws5.max_row, column=1).alignment = WRAP
    ws5.row_dimensions[ws5.max_row].height = 75
    ws5.append([])
ws5.column_dimensions["A"].width = 130

out = os.path.join(HERE, "Equitas_SFB_Financial_Spreads.xlsx")
wb.save(out)
wb2 = openpyxl.load_workbook(out, data_only=True)
cells = sum(1 for s in wb2.sheetnames for row in wb2[s].iter_rows() for c in row if c.value is not None)
nfail_ext = sum(1 for r in wb2["ExternalTieOut"].iter_rows(min_row=4, values_only=True) if r[5] == "CHECK")
print("saved %s; cells=%d; tie-outs %d/%d PASS; external CHECKs=%d" % (os.path.basename(out), cells, npass,
                                                                       len(T), nfail_ext))
