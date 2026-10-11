"""Fills Rating_Scorecard_Equitas Small Finance Bank Limited.xlsx (Step 12b; Stage 2 = Step 17c).

Sector 'Bank / SFB' -> the template's FINANCIAL-SECTOR block (rows 49-55) replaces the corporate
financials. Latest = audited FY2026 from compute_ratios.py (CareEdge definitions). Expected FY+1 =
COPIED (never linked) from Equitas_SFB_Lending_Model_10Oct2026.xlsx base case."""
import importlib.util
import shutil
import sys
from pathlib import Path

import openpyxl

CASE = Path(__file__).resolve().parent
ROOT = CASE.parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from make_spreads import _cache_values, _force_recalc  # noqa: E402

spec = importlib.util.spec_from_file_location("cr", CASE / "compute_ratios.py")
E = importlib.util.module_from_spec(spec)
spec.loader.exec_module(E)
R = E.compute()
F = R["FY2026"]
OUT = CASE / "Rating_Scorecard_Equitas Small Finance Bank Limited.xlsx"
MODEL = CASE / "Equitas_SFB_Lending_Model_10Oct2026.xlsx"


def model(tab, label, col=6):   # col 6 = FY2027E (B..E history, F first projection)
    ws = openpyxl.load_workbook(MODEL, data_only=True)[tab]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label:
            return ws.cell(r, col).value
    raise KeyError(label)


fy1 = dict(crar=model("Credit Ratios", "CRAR - (Tier I + Tier II) / RWA (%)"),
           rota=model("Credit Ratios", "ROTA - PAT / avg total assets (%)"),
           nim=model("Credit Ratios", "NIM - NII / avg total assets (%)"),
           cti=model("Credit Ratios", "Cost to income - opex / (NII + other income) (%)"),
           ta=model("Balance Sheet", "TOTAL ASSETS"))

shutil.copy(ROOT / "Rating_Scorecard.xlsx", OUT)
wb = openpyxl.load_workbook(OUT)
ws = wb["Inputs"]


def put(row, c, d=None, note=None):
    ws.cell(row, 3).value = c
    if d is not None:
        ws.cell(row, 4).value = d
    if note:
        ws.cell(row, 6).value = note


put(4, "Equitas Small Finance Bank Limited")
put(5, "CareEdge")
put(6, "Bank / SFB")
put(7, "10-Oct-2026")
src = "source: forecast model (Equitas_SFB_Lending_Model_10Oct2026.xlsx, base case FY2027E)"
put(49, round(F["CRAR (%)"] / 100, 6), round(fy1["crar"] / 100, 6), "CRAR FY2026 20.31% (Basel II); FY+1 " + src)
put(50, round(F["GNPA (%)"] / 100, 6), round(F["GNPA (%)"] / 100, 6),
    "GNPA FY2026 2.60%; FY+1 held flat - the model projects credit cost, not GNPA (analyst judgment)")
put(51, round(F["NNPA (%)"] / 100, 6), round(F["NNPA (%)"] / 100, 6), "NNPA FY2026 0.72%; FY+1 held flat")
put(52, round(R["FY2026"]["ROTA (%)"] / 100, 6), round(fy1["rota"] / 100, 6),
    "ROTA (CareEdge: PAT / avg total assets) FY2026 0.18%; FY+1 " + src)
put(53, round(R["FY2026"]["NIM (%)"] / 100, 6), round(fy1["nim"] / 100, 6),
    "NIM on avg total assets FY2026 5.98% (bank's own basis 6.57%); FY+1 " + src)
put(54, round(F["Cost to income (%)"] / 100, 6), round(fy1["cti"] / 100, 6), "cost to income FY2026 71.67%")
put(55, round(F["Total assets"], 2), round(fy1["ta"], 2), "total assets Rs crore")
ws["C39"] = 3
ws["F39"] = ("Override 2 -> 3. The 'Bank / SFB' base reflects universal banks; an SFB lends to "
             "self-employed, micro and small borrowers (priority sector 29,698 of 43,580 gross advances), "
             "and the FY2025-26 microfinance stress (Q1 FY2026 loss Rs 223.78 crore) shows the "
             "segment's cyclicality.")
put(40, 3, None, "3 - established, diversified and largely secured franchise (1,060 branches; total assets "
                 "Rs 60,610 crore; microfinance cut to ~10% of advances; top-20 borrowers 3.93%), but CASA "
                 "eroded 42.2% -> 26.2% (FY2023-26) and top-20 depositors rose to 20.61%.")
put(41, 2, None, "2 - board of 11 with 9 independent directors (82%); unmodified audit opinions; no RBI "
                 "penalty in FY2026; frauds small (Rs 18.33 crore). Founder-led management; not 1 given the "
                 "two-year profitability slide and rising credit costs.")
put(42, 0, None, "0 - widely held listed bank with no promoter (holding company amalgamated in FY2023); "
                 "no parent support claimed.")
put(43, "Adequate", None, "LCR 138.45% (Q4 FY2026 avg), NSFR 114.01%; negative cumulative ALM gaps "
                          "within one year offset by excess SLR and RBI facilities - see workpaper s11.")
put(44, "Stable")
# ---- Stage 2: Rating Committee decision, Step 17c (Committee_Minutes_10Oct2026.md) ----
ws["C63"] = 1
ws["D63"] = ("+1 Peer benchmarking / calibration (Committee, 10-Oct-2026): the composite reads the FY2026 "
             "trough ROTA (0.18%) as structural. Capital buffer: CRAR 20.31%, Tier I 16.68%; a FY2027 "
             "credit cost of 3.99% of average assets (2x FY2026) is needed to reach the 15% minimum "
             "(lending model). Recovery already reported: Q4 FY2026 PAT Rs 212.68 crore, Q1 FY2027 PBT "
             "Rs 244.07 crore. De-risked book: microfinance ~10%, 71% CGFMU-covered; top-20 borrowers "
             "3.93%. CareEdge's SFB peer at AA- (Ujjivan: CAR 21.6%, CASA 27%, micro-banking 52%) has a "
             "riskier mix. Not +2 (AA): ROTA below 0.5% for two years, CASA 26%, depositor concentration "
             "20.6%, negative ALM gaps.")
wb.save(OUT)
_force_recalc(OUT)
_cache_values(OUT)
wbv = openpyxl.load_workbook(OUT, data_only=True)
rat, sc = wbv["Rating"], wbv["Scorecard"]
for r in (10, 11, 15, 16, 17, 32, 34, 38):
    print("  Rating!C%-3d %-55s %s" % (r, str(rat.cell(r, 2).value)[:55], rat.cell(r, 3).value))
for r in (39, 40, 41, 42):
    print("  Scorecard!C%-3d %-50s %s" % (r, str(sc.cell(r, 2).value)[:50], sc.cell(r, 3).value))
