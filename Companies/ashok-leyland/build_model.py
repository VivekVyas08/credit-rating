# -*- coding: utf-8 -*-
"""Step 12 for Ashok Leyland (AL standalone - the rated entity; basis per R12, see compute_ratios.py).

Equity-research initiating-coverage Task 2 (model) + Task 3 (valuation tab) + Task 4 (charts),
scaled per case.json step12_scaling_decision. Writes:
  Ashok Leyland_Financial_Model_10Oct2026.xlsx
     Revenue Model | Income Statement | Cash Flow | Balance Sheet | Scenarios | DCF Inputs |
     Credit Ratios (Excel FORMULAS off the three statements) | Debt Breakup | Valuation
  charts/chart_*.png

Historical columns come from compute_ratios.py (R6 - never retyped). Every forward driver is in
ASSUMPTIONS with its source; where it is analyst judgment the label says so. The debt path is
DERIVED from the funding gap (failure mode 18): borrowing happens only when free liquidity would
fall below the floor. Run, then tools/recalc_workbook.py on the output (caches the formulas).
"""
import importlib.util
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("cr", os.path.join(HERE, "compute_ratios.py"))
E = importlib.util.module_from_spec(spec)
spec.loader.exec_module(E)
R = E.compute()
v, tot = E.v, E.tot
H = E.YEARS
P = ["FY2027E", "FY2028E", "FY2029E", "FY2030E", "FY2031E"]
OUT = os.path.join(HERE, "Ashok Leyland_Financial_Model_10Oct2026.xlsx")

PRODUCT_LINES = ["Commercial vehicles", "Engines and gensets", "Ferrous castings and patterns",
                 "Spare parts and others", "Sale of services", "Scrap sales",
                 "Other operating revenue - others"]


# ===================================================================== historical statements
def hist(p):
    s = lambda t, *ks: sum(v(t, k, p) for k in ks)  # noqa: E731
    h = dict(
        vol=v(E.OPS, "Total vehicle volume", p),
        exp_units=v(E.OPS, "Export volume", p),
        rev=v(E.SA_PL, "Revenue from operations", p),
        cogs=E.sa_cogs(p), emp=v(E.SA_PL, "Employee benefits expense", p),
        oth=v(E.SA_PL, "Other expenses", p), oi=v(E.SA_PL, "Other income", p),
        da=v(E.SA_PL, "Depreciation and amortisation", p), fc=v(E.SA_PL, "Finance costs", p),
        exc=v(E.SA_PL, "Exceptional items", p), tax=v(E.SA_PL, "Tax expense", p),
        # balance sheet, grouped
        block=s(E.SA_NCA, "Property, plant and equipment", "Right-of-use asset", "Other intangible assets"),
        cwip=v(E.SA_NCA, "Capital work-in-progress", p),
        iud=v(E.SA_NCA, "Intangibles under development", p),
        gw=v(E.SA_NCA, "Goodwill", p), nci=v(E.SA_NCA, "Non-current investments", p),
        onca=s(E.SA_NCA, "Non-current trade receivables", "Other non-current financial assets",
               "Income tax assets", "Other non-current assets"),
        inv=v(E.SA_CA, "Inventories", p), rec=v(E.SA_CA, "Trade receivables", p),
        free=E.sa_free_liquidity(p),
        earm=E.sa_earmarked_bank(p) + v(E.SA_EARMARKED, "Non-liquid current investments", p),
        oca=s(E.SA_CA, "Current loans", "Other current financial assets", "Other current assets"),
        hfsa=v(E.SA_HFS_ASSETS, "Assets held for sale", p),
        eq=E.sa_equity(p), ncb=v(E.SA_NCL, "Non-current borrowings", p),
        ncl_lease=v(E.SA_NCL, "Non-current lease liabilities", p),
        dtl=v(E.SA_NCL, "Deferred tax liability (net)", p),
        oncl=s(E.SA_NCL, "Other non-current financial liabilities", "Non-current contract liabilities",
               "Non-current provisions"),
        cb=v(E.SA_CL, "Current borrowings", p), cl_lease=v(E.SA_CL, "Current lease liabilities", p),
        pay=s(E.SA_CL, "Trade payables - MSE", "Trade payables - others"),
        ocl=s(E.SA_CL, "Other current financial liabilities", "Current contract liabilities",
              "Current provisions", "Other current liabilities", "Current tax liabilities"),
        hfsl=v(E.SA_HFS_LIAB, "Liabilities held for sale", p),
        # cash flow
        cfo=v(E.SA_CF, "Net cash from operating activities", p),
        capex=v(E.SA_CF, "Capex (purchase of PPE & intangibles)", p),
        support=v(E.SA_CF, "Investment in subsidiaries/JVs", p),
        div=v(E.SA_CF, "Dividend paid", p),
        repay=v(E.SA_CF, "Repayment of non-current borrowings", p),
        newb=v(E.SA_CF, "Proceeds from non-current borrowings", p),
        intpaid=v(E.SA_CF, "Interest paid", p),
    )
    for k in PRODUCT_LINES + ["Rebates and discounts", "Export incentives and grant income"]:
        h[k] = v(E.REV_MIX, k, p)
    h["India"], h["Outside India"] = v(E.REV_GEO, "India", p), v(E.REV_GEO, "Outside India", p)
    return finish(h)


def finish(h):
    h["opbdita"] = h["rev"] - h["cogs"] - h["emp"] - h["oth"]
    h["pbt_pre"] = h["opbdita"] + h["oi"] - h["da"] - h["fc"]
    h["pbt"] = h["pbt_pre"] + h["exc"]
    h["pat"] = h["pbt"] - h["tax"]
    h["ta"] = sum(h[k] for k in ("block", "cwip", "iud", "gw", "nci", "onca", "inv", "rec", "free",
                                 "earm", "oca", "hfsa"))
    h["tle"] = sum(h[k] for k in ("eq", "ncb", "ncl_lease", "dtl", "oncl", "cb", "cl_lease", "pay",
                                  "ocl", "hfsl"))
    return h


HIST = {p: hist(p) for p in H}
for p in H:  # the grouping must reproduce the AR's printed totals before anything is projected
    h = HIST[p]
    assert abs(h["opbdita"] - R[p]["OPBDITA"]) < 0.01, p
    assert abs(h["pat"] - v(E.SA_PL_PRINTED, "Profit for the year", p)) < 0.02, p
    assert abs(h["ta"] - v(E.SA_BS_PRINTED, "Total assets", p)) < 0.02, (p, h["ta"])
    assert abs(h["tle"] - h["ta"]) < 0.02, (p, h["tle"], h["ta"])
F = HIST["FY2026"]

# ===================================================================== assumptions
SHARES = 587.3854552  # crore shares, Rs 1 each - Board's report: Rs 5,87,38,54,552 paid-up
COMMITTED_DIV_FY27 = 1468.46  # 2nd interim Rs 2.50/share declared 28-May-2026 (FY2026 AR 'Dividend')
FLOOR = 1500.0
SCHED_REPAY = [483.82, 129.47, 129.47, 129.47, 129.47]

BASE = dict(
    vol=[5.0, 5.0, 2.0, -4.0, 4.0], real=[1.5] * 5,
    eng=[10.0, 10.0, 8.0, 6.0, 8.0], cast=[3.0] * 5, spares=[8.0, 8.0, 7.0, 5.0, 7.0],
    serv=[3.0] * 5, scrap=[3.0] * 5, othop=[3.0] * 5,
    margin=[13.0, 13.0, 12.5, 11.5, 12.5], emp=[7.0] * 5,
    capex=[1200.0, 1150.0, 1100.0, 1000.0, 1100.0], support=[500.0, 500.0, 400.0, 400.0, 400.0],
    interim=1.0, payout=55.0, exports=[19500, 21000, 22500, 23500, 25000],
    outside=[7.5, 8.0, 8.5, 9.0, 9.5])
BULL = dict(BASE, vol=[8.0, 8.0, 6.0, 4.0, 5.0], real=[2.0] * 5,
            margin=[13.5, 14.0, 14.5, 14.5, 14.5], support=[300.0] * 5,
            exports=[20000, 22500, 25000, 27000, 29000], outside=[7.8, 8.5, 9.2, 9.8, 10.4])
BEAR = dict(BASE, vol=[-8.0, -25.0, -10.0, 15.0, 10.0], real=[0.0, -1.0, 0.0, 1.0, 1.5],
            eng=[3.0, -5.0, 0.0, 5.0, 5.0], spares=[3.0, -5.0, 0.0, 5.0, 5.0],
            margin=[10.5, 6.0, 7.5, 10.0, 11.0], emp=[4.0, 0.0, 0.0, 4.0, 5.0],
            capex=[1000.0, 800.0, 800.0, 900.0, 1000.0],
            support=[1000.0, 1000.0, 1000.0, 500.0, 500.0], interim=0.0, payout=30.0,
            exports=[18000, 15500, 15000, 17000, 19000], outside=[7.0] * 5)
TAIL = dict(BEAR, vol=[-37.0, -20.0, 27.0, 40.0, 10.0], real=[0.0, 0.0, 3.0, 3.0, 2.0],
            eng=[-15.0, -5.0, 15.0, 15.0, 8.0], spares=[-15.0, -5.0, 15.0, 15.0, 8.0],
            margin=[6.0, 3.0, 8.0, 11.0, 12.5], emp=[0.0, 0.0, 3.0, 5.0, 5.0],
            capex=[900.0, 700.0, 800.0, 1000.0, 1100.0],
            exports=[14000, 12000, 14000, 17000, 19000])
SCEN = {"Base": BASE, "Bull": BULL, "Bear / rating stress": BEAR, "Tail (FY2019-21 repeat)": TAIL}

ASSUMPTION_SOURCES = [
    ("Total vehicle volume growth (%)", "vol",
     "ANALYST JUDGMENT. FY2027 per FY2026 AR MD&A: 'FY27 is expected to mark a year of disciplined, "
     "consolidation-led growth'. FY2029-30 moderation reflects the company's own ten-year table "
     "(FY2019 peak 1,97,366 -> FY2021 1,00,725 vehicles): a CV upcycle (FY2022-26) is not assumed "
     "to run unbroken. Timing is illustrative, not a forecast."),
    ("CV realisation per vehicle growth (%)", "real",
     "ANALYST JUDGMENT. History (CV revenue / vehicles): Rs 16.46 lakh FY2023 -> 16.71 lakh FY2026, "
     "about +0.5% a year; 1.5% allows for premiumisation (chairman's letter) and price increases."),
    ("Engines and gensets growth (%)", "eng",
     "ANALYST JUDGMENT, well below the FY2023-26 CAGR of 30.7% (633.84 -> 1,414.29, note 3.7.1); "
     "Power Solutions crossed Rs 1,000 crore 'for the second year' (chairman)."),
    ("Spare parts and others growth (%)", "spares",
     "ANALYST JUDGMENT, below the FY2023-26 CAGR of 20.0% (2,577.88 -> 4,450.37, note 3.7.1)."),
    ("Castings / services / scrap / other operating revenue growth (%)", "serv",
     "ANALYST JUDGMENT - roughly inflation; FY2023-26 services grew 1,106.88 -> 1,257.11."),
    ("OPBDITA margin (%)", "margin",
     "ANALYST JUDGMENT anchored on FY2026 actual 13.03% (engine). Management's 'mid-teen EBITDA in "
     "the medium term' (Q1 FY2025 press release) is used only in the Bull case."),
    ("Employee cost growth (%)", "emp",
     "ANALYST JUDGMENT. FY2023-26: 2,113.86 -> 2,606.69 (7.2% CAGR)."),
    ("Capex (Rs crore)", "capex",
     "ANALYST JUDGMENT. FY2026 actual Rs 1,050 crore (Board's report: Lucknow greenfield bus plant, "
     "alternate-fuel vehicles, digital, plant infrastructure); capital commitments Rs 460.19 crore "
     "(note 3.10)."),
    ("Equity support to subsidiaries (Rs crore)", "support",
     "ANALYST JUDGMENT. FY2024-26 actual 1,567.02 / 717.06 / 386.28 (engine; Optare/Switch, OHM, HLF). "
     "Bear/Tail add a Hinduja Leyland Finance capital call and continued EV losses (ICRA 1006: "
     "'capital requirements for the captive finance business')."),
    ("Export units", "exports",
     "Chairman's letter: 'a clear ambition to reach 25,000 export units in the near term' (FY2026: "
     "18,082). Base reaches 25,000 by FY2031 - slower than the ambition."),
    ("Outside-India share of contract revenue (%)", "outside",
     "FY2026 actual 7.03% (3,090.13 / 43,936.35, note 3.7.1); rises with exports."),
]
FIXED_SOURCES = [
    ("Rebates and discounts (% of gross revenue)", 1.78,
     "FY2026 actual 794.91 / 44,731.26 = 1.78% (note 2.1)."),
    ("Export incentives (% of contract revenue)", 0.16, "FY2026 actual 70.68 / 43,936.35 (note 2.1)."),
    ("Other expenses (% of revenue)", 9.66, "FY2026 actual 4,250.95 / 44,007.03."),
    ("Inventory (% of revenue)", 7.92, "FY2026 actual 3,486.02 / 44,007.03."),
    ("Trade receivables (% of revenue)", 6.70, "FY2026 actual 2,950.41 / 44,007.03."),
    ("Trade payables (% of revenue)", 18.93, "FY2026 actual 8,329.26 / 44,007.03."),
    ("Treasury yield on opening free liquidity (%)", 5.0,
     "FY2026 other income 315.17 / average free liquidity 6,314.66 = 4.99%."),
    ("Interest rate on borrowings (%)", 6.3,
     "FY2026 AR: 'weighted average rate of interest of the borrowing is around 6.30% p.a' "
     "(short-term bank loan note); applied to all borrowings - analyst simplification."),
    ("Non-cash finance cost: unwinding of discount (Rs crore)", 95.69,
     "FY2026 note 2.5; held flat and treated as cash (conservative)."),
    ("Lease interest (Rs crore)", 3.99, "FY2026 note 2.5; leases held flat at Rs 43.34 crore."),
    ("Effective tax rate (%)", 25.5, "FY2026 actual 1,248.49 / 4,814.02 = 25.9%; statutory 25.17%."),
    ("Dividend: committed FY2027 outflow (Rs crore)", COMMITTED_DIV_FY27,
     "2nd interim Rs 2.50/share declared 28-May-2026, outflow Rs 1,468.46 crore (FY2026 AR)."),
    ("Dividend: FY2027 interim (Rs/share) and later payout (% of prior-year PAT)", "1.00 / 55%",
     "FY2026 pattern: Rs 1 interim + Rs 2.50 second interim = 57.7% of FY2026 PAT; Bear/Tail "
     "cut to no interim and 30% payout (policy response, analyst judgment)."),
    ("Scheduled repayments (Rs crore)", "483.82, then 129.47 x4",
     "FY2027 = current maturities 483.82 (note 1.21: term loans 279.25 + NCD 200.00 bullet "
     "17-Mar-2027 + sales-tax 4.91). FY2028-31 = the non-current 517.88 (term loans 456.38 + "
     "sales-tax 61.50) spread evenly - the schedule's own instalments are uneven; approximation."),
    ("Liquidity floor before new borrowing (Rs crore)", FLOOR,
     "ANALYST JUDGMENT: ~2.5 weeks of FY2026 cost of goods sold (31,417.14 / 52 x 2.5). Any "
     "shortfall below it is borrowed (failure mode 18: debt DERIVED from the funding gap)."),
]
OTHER_EXP_PCT, INV_PCT, REC_PCT, PAY_PCT = 9.66, 7.92, 6.70, 18.93
REBATE_PCT, EI_PCT, YIELD, RATE, UNWIND, LEASE_INT, TAX = 1.78, 0.16, 5.0, 6.3, 95.69, 3.99, 25.5
DA_RATE = F["da"] / HIST["FY2025"]["block"]  # FY2026 D&A on opening net block = 13.25%


def project(a):
    rows, prev = [], dict(F)
    prev["real"] = v(E.REV_MIX, "Commercial vehicles", "FY2026") * 100 / F["vol"]
    for i in range(5):
        h = {}
        h["vol"] = prev["vol"] * (1 + a["vol"][i] / 100)
        h["real"] = prev["real"] * (1 + a["real"][i] / 100)
        h["exp_units"] = a["exports"][i]
        h["Commercial vehicles"] = h["vol"] * h["real"] / 100
        for k, key in (("Engines and gensets", "eng"), ("Ferrous castings and patterns", "cast"),
                       ("Spare parts and others", "spares"), ("Sale of services", "serv"),
                       ("Scrap sales", "scrap"), ("Other operating revenue - others", "othop")):
            h[k] = prev[k] * (1 + a[key][i] / 100)
        gross = sum(h[k] for k in PRODUCT_LINES)
        h["Rebates and discounts"] = gross * REBATE_PCT / 100
        rcc = gross - h["Rebates and discounts"]
        h["Export incentives and grant income"] = rcc * EI_PCT / 100
        h["Outside India"] = rcc * a["outside"][i] / 100
        h["India"] = rcc - h["Outside India"]
        h["rev"] = rcc + h["Export incentives and grant income"]
        h["emp"] = prev["emp"] * (1 + a["emp"][i] / 100)
        h["oth"] = h["rev"] * OTHER_EXP_PCT / 100
        opb = h["rev"] * a["margin"][i] / 100
        h["cogs"] = h["rev"] - h["emp"] - h["oth"] - opb
        h["oi"] = prev["free"] * YIELD / 100
        h["da"] = prev["block"] * DA_RATE
        h["exc"] = 0.0
        h["capex"], h["support"] = a["capex"][i], a["support"][i]
        h["repay"] = SCHED_REPAY[i]
        h["div"] = (COMMITTED_DIV_FY27 + a["interim"] * SHARES) if i == 0 else \
            a["payout"] / 100 * max(0.0, prev["pat"])
        # balance sheet movements that do not depend on financing
        for k in ("cwip", "iud", "gw", "onca", "oca", "hfsa", "ncl_lease", "cl_lease", "dtl",
                  "oncl", "ocl", "hfsl"):
            h[k] = prev[k]
        h["block"] = prev["block"] + h["capex"] - h["da"]
        h["nci"] = prev["nci"] + h["support"]
        h["inv"], h["rec"], h["pay"] = (h["rev"] * INV_PCT / 100, h["rev"] * REC_PCT / 100,
                                        h["rev"] * PAY_PCT / 100)
        h["earm"] = prev["earm"] - (100.0 if i == 0 else 0.0)  # DRF deposit applied to the NCD
        # financing: solve new borrowing so free liquidity >= FLOOR (fixed point on interest)
        h["newb"] = 0.0
        for _ in range(30):
            debt_open = prev["ncb"] + prev["cb"]
            debt_close = debt_open - h["repay"] + h["newb"]
            h["fc"] = RATE / 100 * (debt_open + debt_close) / 2 + UNWIND + LEASE_INT
            h["opbdita"] = h["rev"] - h["cogs"] - h["emp"] - h["oth"]
            h["pbt"] = h["pbt_pre"] = h["opbdita"] + h["oi"] - h["da"] - h["fc"]
            h["tax"] = max(0.0, h["pbt"]) * TAX / 100
            h["pat"] = h["pbt"] - h["tax"]
            h["cfo"] = (h["pat"] + h["da"] - (h["inv"] - prev["inv"]) - (h["rec"] - prev["rec"])
                        + (h["pay"] - prev["pay"]))
            pre = (prev["free"] + h["cfo"] - h["capex"] - h["support"] - h["div"] - h["repay"]
                   + (prev["earm"] - h["earm"]))
            need = max(0.0, FLOOR - pre)
            if abs(need - h["newb"]) < 0.005:
                break
            h["newb"] = need
        h["free"] = pre + h["newb"]
        h["intpaid"] = h["fc"]
        # borrowings split: next year's scheduled repayment + Rs 150 crore short-term loans current
        debt_close = prev["ncb"] + prev["cb"] - h["repay"] + h["newb"]
        h["cb"] = 150.0 + (SCHED_REPAY[i + 1] if i < 4 else 0.0)
        h["ncb"] = debt_close - h["cb"]
        h["eq"] = prev["eq"] + h["pat"] - h["div"]
        h = finish(h)
        assert abs(h["ta"] - h["tle"]) < 0.01, ("balance sheet does not balance", i, h["ta"], h["tle"])
        rows.append(h)
        prev = h
    return rows


RUNS = {name: project(a) for name, a in SCEN.items()}
B = RUNS["Base"]


def ratios(h, prev=None):
    td = h["ncb"] + h["cb"] + h["ncl_lease"] + h["cl_lease"]
    tol = h["ta"] - h["eq"]
    ce = td + h["eq"] + h["dtl"] - h["cwip"]
    out = dict(td=td, tnw=h["eq"], tol_tnw=tol / h["eq"], gearing=td / h["eq"],
               td_op=td / h["opbdita"] if h["opbdita"] > 0 else None,
               icr=h["opbdita"] / h["fc"], dscr=(h["pat"] + h["fc"] + h["da"]) / (h["fc"] + h["repay"]),
               margin=100 * h["opbdita"] / h["rev"], net_debt=td - h["free"], free=h["free"],
               nca_td=(h["pat"] + h["da"] - h["div"]) / td,
               cr=(h["inv"] + h["rec"] + h["free"] + h["earm"] + h["oca"])
               / (h["cb"] + h["cl_lease"] + h["pay"] + h["ocl"]))
    if prev is not None:
        pce = (prev["ncb"] + prev["cb"] + prev["ncl_lease"] + prev["cl_lease"] + prev["eq"]
               + prev["dtl"] - prev["cwip"])
        out["roce"] = 100 * (h["opbdita"] - h["da"]) / ((ce + pce) / 2)
    return out


for p, pp in zip(H[1:], H[:-1]):  # Python ratios must reproduce the engine's before projection
    r = ratios(HIST[p], HIST[pp])
    for mine, eng in (("tol_tnw", "TOL/TNW (x)"), ("icr", "Interest coverage (x)"), ("dscr", "DSCR (x)"),
                      ("roce", "RoCE (%)"), ("cr", "Current ratio (x)"), ("td", "Total debt")):
        assert abs(r[mine] - R[p][eng]) < 0.01, (p, mine, r[mine], R[p][eng])

# ===================================================================== workbook
HDR = PatternFill("solid", fgColor="1F3864")
HDRF = Font(bold=True, color="FFFFFF")
SUBF = PatternFill("solid", fgColor="D9E2F3")
PROJF = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
COLS = H + P
ROWS = {}  # (sheet, key) -> row number, for formulas


def sheet(wb, title, first="Rs crore"):
    ws = wb.create_sheet(title)
    ws.append([first, *COLS, "Source / note"])
    for c in ws[1]:
        c.fill, c.font = HDR, HDRF
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions["A"].width = 52
    for i in range(2, 2 + len(COLS)):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.column_dimensions[get_column_letter(2 + len(COLS))].width = 70
    ws.freeze_panes = "B2"
    return ws


def put(ws, key, label, hv, pv, note="", fmt="#,##0.00", bold=False):
    ws.append([label, *[None if x is None else round(x, 4) for x in hv + pv], note])
    r = ws.max_row
    ROWS[(ws.title, key)] = r
    for j in range(2, 2 + len(COLS)):
        c = ws.cell(r, j)
        c.number_format = fmt
        if j >= 2 + len(H):
            c.fill = PROJF
    if bold:
        for c in ws[r]:
            c.font = BOLD
    return r


def section(ws, title):
    ws.append([title])
    ws.cell(ws.max_row, 1).font = BOLD
    for c in ws[ws.max_row]:
        c.fill = SUBF


def series(key, rows=B):
    return [HIST[p][key] for p in H], [r[key] for r in rows]


wb = openpyxl.Workbook()
wb.remove(wb.active)

# ---- Assumptions (cover) ----
ws = wb.create_sheet("Assumptions")
ws.append(["Ashok Leyland Limited - standalone 5-year model (Rs crore). Base case shown on every "
           "statement tab; Scenarios tab carries Bull / Bear (rating stress) / Tail."])
ws.cell(1, 1).font = Font(bold=True, size=12)
ws.append(["Generated by build_model.py from compute_ratios.py (182+ tie-outs). Historical columns "
           "FY2023-FY2026 are audited; yellow = projection. Re-run the script to flex a driver."])
ws.append([])
ws.append(["Driver", *P, "Source / basis"])
for c in ws[ws.max_row]:
    c.fill, c.font = HDR, HDRF
for label, key, src in ASSUMPTION_SOURCES:
    for name, a in SCEN.items():
        ws.append(["%s - %s" % (label, name), *a[key], src if name == "Base" else ""])
ws.append([])
ws.append(["Fixed driver", "Value", "", "", "", "", "Source / basis"])
for c in ws[ws.max_row]:
    c.fill, c.font = HDR, HDRF
for label, val, src in FIXED_SOURCES:
    ws.append([label, val, "", "", "", "", src])
ws.column_dimensions["A"].width = 70
ws.column_dimensions["G"].width = 110
for col in "BCDEF":
    ws.column_dimensions[col].width = 12

# ---- Revenue Model ----
ws = sheet(wb, "Revenue Model")
section(ws, "VOLUMES (units)")
put(ws, "vol", "Total vehicles", *series("vol"), "ten-year table (FY2026 AR p6)", "#,##0")
put(ws, "exp_units", "of which exports", [None] + [HIST[p]["exp_units"] for p in H[1:]],
    [r["exp_units"] for r in B], "chairman's letters; FY2023 not extracted", "#,##0")
put(ws, "real", "CV realisation per vehicle (Rs lakh)",
    [HIST[p]["Commercial vehicles"] * 100 / HIST[p]["vol"] for p in H], [r["real"] for r in B],
    "CV revenue / total vehicles (derived)")
section(ws, "REVENUE BY PRODUCT (note 2.1 / 3.7.1)")
for k in PRODUCT_LINES:
    put(ws, k, k, *series(k), "note 3.7.1 disaggregated revenue")
put(ws, "Rebates and discounts", "Less: rebates and discounts", *series("Rebates and discounts"), "note 2.1")
put(ws, "Export incentives and grant income", "Add: export incentives (+ FY2023 grant income)",
    *series("Export incentives and grant income"), "note 2.1 other operating revenue")
put(ws, "rev", "REVENUE FROM OPERATIONS", *series("rev"), "ties to the P&L (engine tie-outs)", bold=True)
section(ws, "REVENUE BY GEOGRAPHY (contract revenue, note 3.7.1)")
put(ws, "India", "India", *series("India"), "note 3.7.1")
put(ws, "Outside India", "Outside India", *series("Outside India"), "note 3.7.1")
put(ws, "growth", "Revenue growth (%)", [None] + [100 * (HIST[p]["rev"] / HIST[q]["rev"] - 1)
                                                  for p, q in zip(H[1:], H[:-1])],
    [100 * (B[0]["rev"] / F["rev"] - 1)] + [100 * (B[i]["rev"] / B[i - 1]["rev"] - 1) for i in range(1, 5)],
    "derived")

# ---- Income Statement ----
ws = sheet(wb, "Income Statement")
put(ws, "rev", "Revenue from operations", *series("rev"), "P&L", bold=True)
put(ws, "cogs", "Cost of goods sold (materials + purchases + inventory change)", *series("cogs"),
    "P&L; projected = derived residual of the margin assumption")
put(ws, "emp", "Employee benefits expense", *series("emp"), "P&L")
put(ws, "oth", "Other expenses", *series("oth"), "P&L")
put(ws, "opbdita", "OPBDITA (ICRA: excludes all other income)", *series("opbdita"),
    "revenue less operating costs; ties to the company's published EBITDA", bold=True)
put(ws, "oi", "Other income (treasury, dividends, gains)", *series("oi"), "P&L; projected = yield x opening free liquidity")
put(ws, "da", "Depreciation and amortisation", *series("da"), "P&L; projected = FY2026 rate on opening net block")
put(ws, "fc", "Finance costs (incl. unwinding of discount)", *series("fc"), "P&L")
put(ws, "pbt_pre", "Profit before exceptional items and tax", *series("pbt_pre"), "P&L printed subtotal (tied)")
put(ws, "exc", "Exceptional items", *series("exc"), "note 2.8; none projected")
put(ws, "pbt", "Profit before tax", *series("pbt"), "P&L")
put(ws, "tax", "Tax expense", *series("tax"), "P&L")
put(ws, "pat", "PROFIT AFTER TAX", *series("pat"), "P&L printed (tied)", bold=True)

# ---- Cash Flow ----
ws = sheet(wb, "Cash Flow")
put(ws, "cfo", "Net cash from operating activities", *series("cfo"),
    "historical: standalone cash-flow statement; projected: PAT + D&A - change in inventory and "
    "receivables + change in payables (finance cost treated as paid)", bold=True)
put(ws, "capex", "Capex (purchase of PPE and intangibles)", *series("capex"), "cash-flow statement")
put(ws, "support", "Equity investment in subsidiaries / JVs", *series("support"),
    "cash-flow statement + Board's report (engine)")
put(ws, "div", "Dividends paid", *series("div"), "cash-flow statement")
put(ws, "repay", "Repayment of borrowings", *series("repay"), "cash-flow statement; projected = schedule")
put(ws, "newb", "Proceeds from borrowings", *series("newb"),
    "cash-flow statement; projected = DERIVED funding-gap draw (zero unless free liquidity < floor)")
hist_free = [HIST[p]["free"] for p in H]
other = [None] + [(HIST[p]["free"] + HIST[p]["earm"]) - (HIST[q]["free"] + HIST[q]["earm"])
                  - (HIST[p]["cfo"] - HIST[p]["capex"] - HIST[p]["support"] - HIST[p]["div"]
                     - HIST[p]["repay"] + HIST[p]["newb"]) for p, q in zip(H[1:], H[:-1])]
put(ws, "other", "Other flows, net (interest paid, leases, treasury, investment sales)",
    other, [pr["earm"] - r["earm"] for r, pr in zip(B, [F] + B[:-1])],
    "historical = balancing item to the movement in cash + bank + current investments; "
    "projected = movement in earmarked balances only")
put(ws, "free", "Closing free cash and liquid investments", hist_free, [r["free"] for r in B],
    "unencumbered: excludes unclaimed-dividend, escrow, DRF and FY2024 dividend accounts", bold=True)

# ---- Balance Sheet ----
ws = sheet(wb, "Balance Sheet")
BS = [("block", "Net block (PPE + right-of-use + other intangibles)"), ("cwip", "Capital work-in-progress"),
      ("iud", "Intangible assets under development"), ("gw", "Goodwill"),
      ("nci", "Non-current investments (incl. subsidiaries)"),
      ("onca", "Other non-current assets"), ("inv", "Inventories"), ("rec", "Trade receivables"),
      ("free", "Free cash and liquid investments"),
      ("earm", "Earmarked bank balances and non-liquid investments"),
      ("oca", "Other current assets"), ("hfsa", "Assets held for sale"), ("ta", "TOTAL ASSETS"),
      ("eq", "Equity (= tangible net worth; no revaluation reserve)"),
      ("ncb", "Non-current borrowings"), ("ncl_lease", "Non-current lease liabilities"),
      ("dtl", "Deferred tax liability (net)"), ("oncl", "Other non-current liabilities and provisions"),
      ("cb", "Current borrowings (incl. current maturities)"), ("cl_lease", "Current lease liabilities"),
      ("pay", "Trade payables"), ("ocl", "Other current liabilities and provisions"),
      ("hfsl", "Liabilities held for sale"), ("tle", "TOTAL EQUITY AND LIABILITIES")]
for k, lab in BS:
    put(ws, k, lab, *series(k), "standalone balance sheet (grouped; totals tie to the AR)" if k in ("ta", "tle") else "",
        bold=k in ("ta", "tle", "eq"))
put(ws, "chk", "Check: assets - (equity + liabilities)", [HIST[p]["ta"] - HIST[p]["tle"] for p in H],
    [r["ta"] - r["tle"] for r in B], "must be zero in every column")

# ---- Credit Ratios: FORMULAS off the three statements ----
ws = sheet(wb, "Credit Ratios", "ICRA convention (formulas)")


def ref(tab, key, col):
    return "'%s'!%s%d" % (tab, col, ROWS[(tab, key)])


IS_, BS_, CF_ = "Income Statement", "Balance Sheet", "Cash Flow"
for key, lab, fmt, f in (
    ("m", "OPBDITA margin (%)", "0.00", lambda c, pc: "=100*%s/%s" % (ref(IS_, "opbdita", c), ref(IS_, "rev", c))),
    ("pm", "PAT margin (%)", "0.00", lambda c, pc: "=100*%s/%s" % (ref(IS_, "pat", c), ref(IS_, "rev", c))),
    ("td", "Total debt (borrowings + leases)", "#,##0.00", lambda c, pc: "=%s+%s+%s+%s" % (
        ref(BS_, "ncb", c), ref(BS_, "cb", c), ref(BS_, "ncl_lease", c), ref(BS_, "cl_lease", c))),
    ("tnw", "Tangible net worth", "#,##0.00", lambda c, pc: "=%s" % ref(BS_, "eq", c)),
    ("g", "Gearing - total debt / TNW (x)", "0.00", lambda c, pc: "=%s%d/%s%d" % (c, ROWS[("Credit Ratios", "td")], c, ROWS[("Credit Ratios", "tnw")])),
    ("tol", "TOL / TNW (x)", "0.00", lambda c, pc: "=(%s-%s)/%s" % (ref(BS_, "ta", c), ref(BS_, "eq", c), ref(BS_, "eq", c))),
    ("tdop", "Total debt / OPBDITA (x)", "0.00", lambda c, pc: "=%s%d/%s" % (c, ROWS[("Credit Ratios", "td")], ref(IS_, "opbdita", c))),
    ("icr", "Interest coverage - OPBDITA / finance costs (x)", "0.00", lambda c, pc: "=%s/%s" % (ref(IS_, "opbdita", c), ref(IS_, "fc", c))),
    ("dscr", "DSCR - (PAT + interest + D&A) / (interest + repayments) (x)", "0.00", lambda c, pc: "=(%s+%s+%s)/(%s+%s)" % (
        ref(IS_, "pat", c), ref(IS_, "fc", c), ref(IS_, "da", c), ref(IS_, "fc", c), ref(CF_, "repay", c))),
    ("nca", "NCA / total debt (x)", "0.00", lambda c, pc: "=(%s+%s-%s)/%s%d" % (
        ref(IS_, "pat", c), ref(IS_, "da", c), ref(CF_, "div", c), c, ROWS[("Credit Ratios", "td")])),
    ("roce", "RoCE - PBIT / avg (debt + TNW + DTL - CWIP) (%)", "0.00", lambda c, pc: None if pc is None else
        "=100*(%s-%s)/((%s%d+%s+%s-%s+%s%d+%s+%s-%s)/2)" % (
            ref(IS_, "opbdita", c), ref(IS_, "da", c), c, ROWS[("Credit Ratios", "td")], ref(BS_, "eq", c),
            ref(BS_, "dtl", c), ref(BS_, "cwip", c), pc, ROWS[("Credit Ratios", "td")], ref(BS_, "eq", pc),
            ref(BS_, "dtl", pc), ref(BS_, "cwip", pc))),
    ("cr", "Current ratio (x)", "0.00", lambda c, pc: "=(%s+%s+%s+%s+%s)/(%s+%s+%s+%s)" % (
        ref(BS_, "inv", c), ref(BS_, "rec", c), ref(BS_, "free", c), ref(BS_, "earm", c), ref(BS_, "oca", c),
        ref(BS_, "cb", c), ref(BS_, "cl_lease", c), ref(BS_, "pay", c), ref(BS_, "ocl", c))),
    ("nd", "Net debt (negative = net cash)", "#,##0.00", lambda c, pc: "=%s%d-%s" % (c, ROWS[("Credit Ratios", "td")], ref(BS_, "free", c))),
):
    ws.append([lab])
    r = ws.max_row
    ROWS[("Credit Ratios", key)] = r
    for j in range(len(COLS)):
        c = get_column_letter(2 + j)
        pc = get_column_letter(1 + j) if j > 0 else None
        cell = ws.cell(r, 2 + j)
        cell.value = f(c, pc)
        cell.number_format = fmt
        if j >= len(H):
            cell.fill = PROJF
ws.append([])
ws.append(["ICRA indicative bands (960 / 1006): RoCE >=25% strongest, <10% weakest; TOL/TNW <=0.9x / "
           ">3.0x; Total debt/OPBDITA <=0.5x / >5.0x; Interest coverage >=18x / <2x; DSCR >=4.0x / <1.1x."])

# ---- Scenarios ----
ws = wb.create_sheet("Scenarios")
ws.append(["Scenario / metric", *P, "Note"])
for c in ws[1]:
    c.fill, c.font = HDR, HDRF
ws.column_dimensions["A"].width = 58
ws.column_dimensions["G"].width = 60
SCEN_OUT = {}
for name, rows in RUNS.items():
    section(ws, name.upper())
    prevs = [F] + rows[:-1]
    rr = [ratios(h, pv) for h, pv in zip(rows, prevs)]
    SCEN_OUT[name] = rr
    for lab, f, fmt in (("Total vehicles", lambda h, r: h["vol"], "#,##0"),
                        ("Revenue from operations", lambda h, r: h["rev"], "#,##0"),
                        ("OPBDITA margin (%)", lambda h, r: r["margin"], "0.00"),
                        ("PAT", lambda h, r: h["pat"], "#,##0"),
                        ("Free cash and liquid investments", lambda h, r: r["free"], "#,##0"),
                        ("New borrowing drawn (funding gap)", lambda h, r: h["newb"], "#,##0"),
                        ("Total debt", lambda h, r: r["td"], "#,##0"),
                        ("Net debt (negative = net cash)", lambda h, r: r["net_debt"], "#,##0"),
                        ("TOL / TNW (x)", lambda h, r: r["tol_tnw"], "0.00"),
                        ("Total debt / OPBDITA (x)", lambda h, r: r["td_op"], "0.00"),
                        ("Interest coverage (x)", lambda h, r: r["icr"], "0.00"),
                        ("DSCR (x)", lambda h, r: r["dscr"], "0.00"),
                        ("RoCE (%)", lambda h, r: r["roce"], "0.00")):
        ws.append([lab, *[None if f(h, r) is None else round(f(h, r), 2) for h, r in zip(rows, rr)]])
        for j in range(2, 7):
            ws.cell(ws.max_row, j).number_format = fmt

# ---- DCF Inputs ----
ws = sheet(wb, "DCF Inputs")
ufcf_h = [None] * 4
nwc = lambda h: h["inv"] + h["rec"] - h["pay"]  # noqa: E731
ufcf = []
for i, h in enumerate(B):
    pv_ = F if i == 0 else B[i - 1]
    ufcf.append(h["opbdita"] - TAX / 100 * (h["opbdita"] - h["da"]) - h["capex"] - (nwc(h) - nwc(pv_)))
put(ws, "opbdita", "OPBDITA", *series("opbdita"), "Income Statement")
put(ws, "taxop", "Less: tax on operating profit (25.5% x (OPBDITA - D&A))",
    [None] * 4, [TAX / 100 * (h["opbdita"] - h["da"]) for h in B], "")
put(ws, "capex", "Less: capex", *series("capex"), "")
put(ws, "dnwc", "Less: increase in operating working capital (inventory + receivables - payables)",
    [None] * 4, [nwc(h) - nwc(F if i == 0 else B[i - 1]) for i, h in enumerate(B)], "")
put(ws, "ufcf", "UNLEVERED FREE CASH FLOW (standalone, before support to subsidiaries)",
    ufcf_h, ufcf, "", bold=True)

# ---- Valuation (Task 3: role = financial flexibility only, never the grade) ----
WACC, G = 12.5, 5.0
def dcf(wacc, g):  # noqa: E302
    pv = sum(cf / (1 + wacc / 100) ** (i + 1) for i, cf in enumerate(ufcf))
    tv = ufcf[-1] * (1 + g / 100) / ((wacc - g) / 100)
    return pv + tv / (1 + wacc / 100) ** 5


ws = wb.create_sheet("Valuation")
ws.append(["DCF of the standalone CV business (Rs crore) - informs financial flexibility only "
           "(Step 12(f)); it does not move the grade."])
ws.cell(1, 1).font = BOLD
ws.append(["Cost of equity = risk-free 6.5% + beta 1.1 x ERP 5.5% = 12.55%; AL is net-cash, so "
           "WACC ~= cost of equity, rounded to 12.5%. Terminal growth 5.0% (nominal). ALL ANALYST "
           "ASSUMPTIONS - not market-sourced."])
EV = dcf(WACC, G)
ws.append(["Enterprise value (WACC 12.5%, g 5.0%)", round(EV, 0)])
ws.append(["Add: free cash and liquid investments, 31-Mar-2026", round(F["free"], 0)])
ws.append(["Less: total debt, 31-Mar-2026", round(R["FY2026"]["Total debt"], 0)])
ws.append(["Add: investments in subsidiaries/associates at book (incl. HLF), 31-Mar-2026",
           round(F["nci"], 0), "book value; HLF is unlisted - no market value available"])
EQV = EV + F["free"] - R["FY2026"]["Total debt"] + F["nci"]
ws.append(["Implied equity value (Rs crore)", round(EQV, 0)])
ws.append(["Implied value per share (Rs)", round(EQV / SHARES, 1), "587.39 crore shares"])
PRICE = 150.48  # NSE close 01-Oct-2026 per incredmoney.com aggregator - NOT verified against NSE
MCAP = PRICE * SHARES
ws.append([])
ws.append(["Market price (Rs, 01-Oct-2026; third-party aggregator, not NSE-verified)", PRICE,
           "52-week range Rs 134.21-215.42 (same source)"])
ws.append(["Market capitalisation (Rs crore) = price x 587.39 crore shares", round(MCAP, 0)])
ws.append(["Market capitalisation / standalone total debt 31-Mar-2026 (x)",
           round(MCAP / R["FY2026"]["Total debt"], 1),
           "the financial-flexibility reading: equity cushion over debt"])
ws.append(["DCF value / market price", round(EQV / SHARES / PRICE, 2),
           "below 1: the market prices more growth or a lower discount rate than this conservative DCF"])
ws.append([])
ws.append(["Sensitivity: equity value (Rs crore)", *["g %.1f%%" % g for g in (4.0, 4.5, 5.0, 5.5, 6.0)]])
SENS = {}
for w in (11.0, 11.75, 12.5, 13.25, 14.0):
    row = [dcf(w, g) + F["free"] - R["FY2026"]["Total debt"] + F["nci"] for g in (4.0, 4.5, 5.0, 5.5, 6.0)]
    SENS[w] = row
    ws.append(["WACC %.2f%%" % w, *[round(x, 0) for x in row]])
ws.column_dimensions["A"].width = 78
for col in "BCDEF":
    ws.column_dimensions[col].width = 13

# ---- Debt Breakup ----
ws = wb.create_sheet("Debt Breakup")
ws.column_dimensions["A"].width = 66
for col in "BCD":
    ws.column_dimensions[col].width = 13
ws.column_dimensions["E"].width = 90


def drow(*vals):
    ws.append(list(vals))


drow("(i) STATUTORY BORROWINGS - standalone (Rs crore)", "FY2024", "FY2025", "FY2026", "Source")
for c in ws[1]:
    c.fill, c.font = HDR, HDRF
for k, lab in (("Non-current borrowings", "Non-current borrowings"), ("Current borrowings", "Current borrowings (incl. current maturities)")):
    t = E.SA_NCL if k.startswith("Non") else E.SA_CL
    drow(lab, *[v(t, k, p) for p in H[1:]], "balance sheet")
drow("Lease liabilities (non-current + current)",
     *[v(E.SA_NCL, "Non-current lease liabilities", p) + v(E.SA_CL, "Current lease liabilities", p) for p in H[1:]],
     "balance sheet")
drow("TOTAL DEBT", *[round(R[p]["Total debt"], 2) for p in H[1:]],
     "sum of the three lines; cross-checked to screener.in in the spreads' ExternalTieOut tab")
drow()
drow("(ii) INSTRUMENT VIEW at 31-Mar-2026 (face value, FY2026 AR repayment schedule)", "", "", "FY2026", "")
for k in ("Term loans - non-current", "Term loans - current maturities", "NCD Series 3 7.30% AL 2027 - current",
          "Sales-tax loans (interest-free) - non-current", "Sales-tax loans (interest-free) - current",
          "Short-term loans from banks (unsecured)"):
    note = {"NCD Series 3 7.30% AL 2027 - current":
            "ISIN INE208A07406; bullet 17-Mar-2027; call option end of year 3; secured 1.10x on Hosur P&M",
            "Short-term loans from banks (unsecured)": "note 1.21"}.get(k, "borrowings note repayment schedule")
    drow(k, "", "", v(E.SA_DEBT_FY26, k, "FY2026"), note)
drow("Commercial paper outstanding", "", "", 0.00, "nil at the standalone level (Rs 2,000 crore rated programme)")
INSTR = ("Term loans - non-current", "Term loans - current maturities", "NCD Series 3 7.30% AL 2027 - current",
         "Sales-tax loans (interest-free) - non-current", "Sales-tax loans (interest-free) - current",
         "Short-term loans from banks (unsecured)")
drow("Sum of instruments (face)", "", "", round(sum(v(E.SA_DEBT_FY26, k, "FY2026") for k in INSTR), 2),
     "vs balance-sheet borrowings 1,151.70: 0.34 effective-interest residual (engine tie-out)")
drow("Currency mix", "", "", "", "no foreign-currency borrowing disclosed at the standalone level; acceptances: none shown separately")
drow()
drow("(iii) THE AR'S OWN NET-DEBT BRIDGE (standalone note, Rs crore)", "FY2024", "FY2025", "FY2026", "")
for k in ("Cash", "Liquid investments", "Current borrowings (bridge)", "Non-current borrowings (bridge)",
          "Derivative asset/(liability)", "Lease liabilities (bridge)", "Net surplus (printed)"):
    drow(k, *[v(E.AR_NET_DEBT_BRIDGE, k, p) for p in H[1:]], "AR net-debt reconciliation")
drow("Arithmetic check: cash + liquid inv - borrowings + derivative - leases",
     *[round(v(E.AR_NET_DEBT_BRIDGE, "Cash", p) + v(E.AR_NET_DEBT_BRIDGE, "Liquid investments", p)
             - v(E.AR_NET_DEBT_BRIDGE, "Current borrowings (bridge)", p)
             - v(E.AR_NET_DEBT_BRIDGE, "Non-current borrowings (bridge)", p)
             + v(E.AR_NET_DEBT_BRIDGE, "Derivative asset/(liability)", p)
             - v(E.AR_NET_DEBT_BRIDGE, "Lease liabilities (bridge)", p), 2) for p in H[1:]],
     "equals the printed net surplus in every year")
drow("Analyst net debt on free liquidity (negative = net cash)", *[round(R[p]["Net debt (negative = net cash)"], 2) for p in H[1:]],
     "basis difference: ours ADDS the free 3-12 month deposits (Rs 800 crore FY2026) and EXCLUDES the "
     "derivative; the bridge excludes all bank deposits. Chairman's 'net cash Rs 5,899 crore' adds every "
     "bank balance incl. earmarked ones.")
drow()
drow("(iv) RATED LIMITS vs DRAWN DEBT", "", "", "Rs crore", "")
drow("ICRA rated programme - CRA-PUBLISHED 02-Jun-2026, NOT issuer-confirmed", "", "", 7479.38,
     "CP 2,000 + NCD 200 + proposed NCD 300 + LT/ST FB 2,000 + LT/ST NFB 1,500 + term loans 729.38 + "
     "LT/ST unallocated 750 (line arithmetic verified)")
drow("Drawn borrowings + leases, 31-Mar-2026", "", "", round(R["FY2026"]["Total debt"], 2),
     "limits and drawn debt are different measures and are not expected to reconcile")
drow("Drawn term loans (face) vs rated term loans 729.38", "", "",
     round(v(E.SA_DEBT_FY26, "Term loans - non-current", "FY2026")
           + v(E.SA_DEBT_FY26, "Term loans - current maturities", "FY2026"), 2),
     "6.25 gap = size of an April-May 2026 instalment; to be confirmed by issuer")
drow("Undrawn bank facilities, 31-Mar-2026", "", "", 2854.34, "liquidity-risk note (secured 2,000.00 + unsecured 854.34)")
for row in ws.iter_rows(min_row=2):
    if row[0].value and str(row[0].value).startswith("("):
        for c in row:
            c.fill, c.font = HDR, HDRF

wb.move_sheet("Assumptions", offset=-len(wb.sheetnames))
wb.save(OUT)
print("wrote", os.path.basename(OUT), "sheets:", wb.sheetnames)

# ===================================================================== console summary
for name, rr in SCEN_OUT.items():
    rows = RUNS[name]
    worst = min(range(5), key=lambda i: rr[i]["icr"])
    print("%-26s FY27 rev %8.0f  trough %s: margin %5.2f%% ICR %6.2fx DSCR %5.2fx TD/OP %s "
          "free %7.0f newdebt %6.0f | FY31 free %7.0f TD %6.0f"
          % (name, rows[0]["rev"], P[worst], rr[worst]["margin"], rr[worst]["icr"], rr[worst]["dscr"],
             "n/m" if rr[worst]["td_op"] is None else "%.2fx" % rr[worst]["td_op"], rr[worst]["free"],
             sum(h["newb"] for h in rows), rr[-1]["free"], rr[-1]["td"]))
print("DCF equity value Rs %.0f crore (Rs %.1f/share) at WACC %.1f%% g %.1f%%" % (EQV, EQV / SHARES, WACC, G))

# ===================================================================== Task 4 - charts
import zipfile  # noqa: E402

import matplotlib  # noqa: E402
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

plt.rcParams.update({"font.family": "Times New Roman", "font.size": 10})
CH = os.path.join(HERE, "charts")
os.makedirs(CH, exist_ok=True)
NAVY, TEAL, GREY, RED = "#1F3864", "#2E8B8B", "#9E9E9E", "#C0392B"
YRS = [p.replace("FY20", "FY") for p in H] + [p.replace("FY20", "FY") for p in P]
# FY2026 AR ten-year table 'A historical perspective' (p6) - vehicles and PAT, FY2017-FY2026
TEN_Y = ["FY17", "FY18", "FY19", "FY20", "FY21", "FY22", "FY23", "FY24", "FY25", "FY26"]
TEN_VOL = [145066, 174873, 197366, 125200, 100725, 128326, 192205, 194555, 195097, 220437]
TEN_PAT = [1223, 1718, 1983, 240, -314, 542, 1380, 2618, 3303, 3566]
INDEX = []


def save(fig, name, title, src):
    path = os.path.join(CH, name)
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    INDEX.append("%s | %s | %s" % (name, title, src))


fig, ax = plt.subplots(figsize=(8, 4))
ax.bar(TEN_Y, [x / 1000 for x in TEN_VOL], color=NAVY)
ax.set_ylabel("Vehicles ('000)")
ax2 = ax.twinx()
ax2.plot(TEN_Y, TEN_PAT, color=RED, marker="o")
ax2.axhline(0, color=GREY, lw=0.6)
ax2.set_ylabel("PAT (Rs crore)")
ax.set_title("Ten years of cyclicality: volumes fell 49% FY19-FY21 and AL made a loss")
save(fig, "chart_01_cyclicality_volumes_pat.png", "Vehicles and PAT FY2017-26", "FY2026 AR ten-year table")

fig, ax = plt.subplots(figsize=(8, 4.2))
stack = [[HIST[p][k] for p in H] + [r[k] for r in B] for k in PRODUCT_LINES]
labels = ["CVs", "Engines & gensets", "Castings", "Spares & others", "Services", "Scrap", "Other op. rev."]
ax.stackplot(YRS, stack, labels=labels, colors=[NAVY, TEAL, "#7FB3D5", "#F5B041", GREY, "#D7DBDD", "#AAB7B8"])
ax.axvline(3.5, color="black", ls="--", lw=0.8)
ax.set_ylabel("Rs crore")
ax.legend(fontsize=7, loc="upper left", ncol=2)
ax.set_title("Revenue by product (gross, before rebates): FY2023-26 actual, FY2027-31 base case")
save(fig, "chart_03_revenue_by_product.png", "Revenue by product", "note 3.7.1; model base case")

fig, ax = plt.subplots(figsize=(8, 4))
ind = [HIST[p]["India"] for p in H] + [r["India"] for r in B]
out = [HIST[p]["Outside India"] for p in H] + [r["Outside India"] for r in B]
ax.bar(YRS, ind, color=NAVY, label="India")
ax.bar(YRS, out, bottom=ind, color=TEAL, label="Outside India")
ax.legend()
ax.set_ylabel("Rs crore")
ax.set_title("Contract revenue by geography (exports 7.0% in FY2026)")
save(fig, "chart_04_revenue_by_geography.png", "Revenue by geography", "note 3.7.1; model base case")

CRH = {p: ratios(HIST[p], HIST[q]) for p, q in zip(H[1:], H[:-1])}
CRH["FY2023"] = ratios(HIST["FY2023"])
CRP = SCEN_OUT["Base"]
fig, ax = plt.subplots(figsize=(8, 4))
ax.plot(YRS, [CRH[p]["margin"] for p in H] + [r["margin"] for r in CRP], color=NAVY, marker="o", label="OPBDITA margin (%)")
ax.set_ylabel("OPBDITA margin (%)")
ax2 = ax.twinx()
ax2.bar(YRS, [CRH[p]["icr"] for p in H] + [r["icr"] for r in CRP], color=TEAL, alpha=0.35, label="Interest coverage (x)")
ax2.axhline(18, color=RED, ls=":", lw=1)
ax2.text(0, 19, "ICRA strongest band >=18x", color=RED, fontsize=8)
ax2.set_ylabel("Interest coverage (x)")
ax.set_title("Margin and coverage: actual FY2023-26, base case FY2027-31")
save(fig, "chart_10_margin_coverage.png", "OPBDITA margin and interest coverage", "engine; model")

fig, ax = plt.subplots(figsize=(8, 4))
ax.bar(YRS, [-CRH[p]["net_debt"] for p in H] + [-r["net_debt"] for r in CRP], color=TEAL, label="Net cash (Rs crore)")
ax.set_ylabel("Net cash (Rs crore)")
ax2 = ax.twinx()
ax2.plot(YRS, [CRH[p]["tol_tnw"] for p in H] + [r["tol_tnw"] for r in CRP], color=NAVY, marker="o", label="TOL/TNW (x)")
ax2.axhline(0.9, color=RED, ls=":", lw=1)
ax2.set_ylabel("TOL/TNW (x)")
ax.set_title("Standalone net cash and TOL/TNW (ICRA strongest band <=0.9x)")
save(fig, "chart_12_net_cash_leverage.png", "Net cash and TOL/TNW", "engine; model")

fig, ax = plt.subplots(figsize=(8, 4))
for (name, rr), col in zip(SCEN_OUT.items(), [NAVY, TEAL, "#F5B041", RED]):
    ax.plot([p.replace("FY20", "FY") for p in P], [r["free"] for r in rr], marker="o", color=col, label=name)
ax.axhline(FLOOR, color=GREY, ls="--", lw=0.8)
ax.text(0, FLOOR + 150, "liquidity floor Rs 1,500 crore", fontsize=8, color=GREY)
ax.set_ylabel("Free cash and liquid investments (Rs crore)")
ax.legend(fontsize=8)
ax.set_title("Free liquidity by scenario - no scenario needs new borrowing")
save(fig, "chart_14_scenario_liquidity.png", "Scenario free liquidity", "model Scenarios tab")

fig, ax = plt.subplots(figsize=(6.5, 4))
gs = [4.0, 4.5, 5.0, 5.5, 6.0]
ws_ = sorted(SENS)
mat = [[SENS[w][j] / SHARES for j in range(5)] for w in ws_]
im = ax.imshow(mat, cmap="RdYlGn", aspect="auto")
ax.set_xticks(range(5), ["g %.1f%%" % g for g in gs])
ax.set_yticks(range(5), ["WACC %.2f%%" % w for w in ws_])
for i in range(5):
    for j in range(5):
        ax.text(j, i, "%.0f" % mat[i][j], ha="center", va="center", fontsize=8)
fig.colorbar(im, label="Rs per share")
ax.set_title("DCF sensitivity - value per share (Rs)")
save(fig, "chart_28_dcf_sensitivity.png", "DCF sensitivity heatmap", "model Valuation tab")

fig, ax = plt.subplots(figsize=(8, 3))
bars = [("DCF (WACC 11-14%, g 4-6%)", min(min(r) for r in SENS.values()) / SHARES, max(max(r) for r in SENS.values()) / SHARES),
        ("52-week trading range", 134.21, 215.42)]
for i, (lab, lo, hi) in enumerate(bars):
    ax.barh(i, hi - lo, left=lo, color=[NAVY, TEAL][i])
    ax.text(hi + 2, i, "Rs %.0f-%.0f" % (lo, hi), va="center", fontsize=8)
ax.axvline(PRICE, color=RED, ls="--")
ax.text(PRICE + 1, 1.45, "Rs %.2f (01-Oct-2026)" % PRICE, color=RED, fontsize=8)
ax.set_yticks(range(len(bars)), [b[0] for b in bars])
ax.set_xlabel("Rs per share")
ax.set_title("Valuation football field (flexibility context only)")
save(fig, "chart_32_valuation_football_field.png", "Valuation football field", "model Valuation tab; aggregator price")

ZIP = os.path.join(HERE, "Ashok Leyland_Charts_10Oct2026.zip")
with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
    for line in INDEX:
        z.write(os.path.join(CH, line.split(" | ")[0]), line.split(" | ")[0])
    z.writestr("chart_index.txt", "file | title | source\n" + "\n".join(INDEX) + "\n")
print("wrote %d charts + %s" % (len(INDEX), os.path.basename(ZIP)))
