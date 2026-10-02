"""Fills Rating_Scorecard_Repco Home Finance Limited.xlsx (CLAUDE.md Step 12b).

Uses the TEMPLATE'S OWN Inputs -> Scorecard -> Rating mechanism (sector = 'Housing Finance
Co (HFC)'), rather than a bespoke sheet, so the anchor grade is the template's own formula
chain reading back its own output - not a parallel calculation.

Financial-sector metrics: "Latest" = FY2025 (the last period with a full, audited balance
sheet - FY2026 carries a P&L but no balance sheet anywhere in the document set, per
compute_ratios.py's BS_NOTE). "Expected FY+1" = FY2027E from the lending model
(Repco Home Finance Limited_Lending_Model.xlsx), which opens on the 30-Jun-2026 actuals.
Every value is copied, never retyped from memory - CLAUDE.md Step 12: "copied (never linked)
into the Step 12b scorecard's assumption cells, each with a 'source: forecast model' note."
"""
import shutil
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
# make_spreads.py moved to tools/ when the project root was decluttered (01-Sep-2026).
sys.path.insert(0, str(ROOT / "tools"))
from make_spreads import _cache_values, _force_recalc          # noqa: E402

SRC = ROOT / "Rating_Scorecard.xlsx"
OUT = Path(__file__).resolve().parent / "Rating_Scorecard_Repco Home Finance Limited.xlsx"


def main():
    shutil.copy(SRC, OUT)
    wb = openpyxl.load_workbook(OUT)
    ws = wb["Inputs"]

    ws["C4"] = "Repco Home Finance Limited"
    ws["C5"] = "ICRA"
    ws["C6"] = "Housing Finance Co (HFC)"
    ws["C7"] = "24-Aug-2026"

    # Corporate financials/projections (rows 12-35) are IGNORED for an HFC sector per the
    # template's own README ("leave these at their example values"); left untouched.

    ws["B39"] = ("Industry risk override (blank = use sector base 2.5 from Ref_Weights) - "
                 "left BLANK. India's HFC sector is well-regulated (NHB/RBI), growing at a "
                 "double-digit industry CAGR per the FY2025 AR MD&A, and the 2.5 base score "
                 "the template assigns HFCs (between Bank 2.0 and NBFC 3.0) is accepted "
                 "as-is.")
    ws["C40"] = 4
    ws["E40"] = ("Franchise and scale - SAME score and rationale as the RatingScore tab of "
                 "the Financial_Spreads workbook: a 25-year-old HFC with a defensible "
                 "tier-2/tier-3 south India niche, but small (AUM Rs 14,492 cr FY2025 vs "
                 "PNB Housing Rs 98,000 cr) and slow-growing (7.2% AUM CAGR FY2022-FY2025).")
    ws["C41"] = 2
    ws["E41"] = ("Unmodified audit opinion FY2023-FY2026 in every filing reviewed; MD "
                 "T Karunakaran has 33 years in housing finance and has been with RHFL since "
                 "2004; board of 8 (5 independent, 2 non-independent non-executive, 1 MD). "
                 "The only qualifier is the Section 36(1)(viii) deferred-tax-reversal "
                 "Emphasis of Matter (Note 7, Q1FY2027 filing) - a disclosed accounting "
                 "judgement, not a governance lapse.")
    ws["C42"] = 0
    ws["E42"] = ("No parent support notch. Repco Bank Ltd holds 37.13% (FY2025 AR) but is a "
                 "promoter shareholding, not a support undertaking: AR General Shareholder "
                 "Information item 47 ('value of obligations...for which a guarantee, surety, "
                 "indemnity or comfort letter has been provided') states NA, and searches for "
                 "'letter of comfort', 'letter of support' and 'guarantee...Repco Bank' return "
                 "no hits. Both ICRA's and CareEdge's existing ratings on this issuer (AR Note "
                 "32.17; CareEdge press release 30-Sep-2025) are on a STANDALONE analytical "
                 "approach, consistent with 0.")
    ws["C43"] = "Adequate"
    ws["E43"] = ("Not Strong: LCR of 202.13% (Q1FY2027) is comfortable, but 12 counterparties "
                 "fund 96.69% of total liabilities and the top 10 borrowings fund 93.34% of "
                 "total borrowings (AR Note 34, 31-Mar-2025) - a single-channel funding base "
                 "with no CP, ECB, secured NCD or public-deposit access as at that date.")
    ws["C44"] = "Stable"

    # ---- Financial-sector inputs (rows 49-55): FY2025 audited "Latest", FY2027E model "Expected"
    rows = [
        (49, "CRAR", 0.3709,
         "FY2025 AR Note 32.1: CRAR 37.09% (Tier I 36.41%, Tier II 0.68%)", 0.4018,
         "Lending_Model Credit_Ratios!C - FY2027E indicative CRAR (net worth less "
         "non-earning assets, over RWA at the FY2025 risk-weight density); labelled "
         "'indicative only' in the model because the regulatory computation deducts items "
         "this model cannot see"),
        (50, "GNPA", 0.0326,
         "FY2025 AR: gross NPA Rs 472.91 cr / 3.26% of loans (excl. interest accrued and EIR; "
         "the RBI sectoral-exposure format gives 3.99% including them - both bases disclosed "
         "in compute_ratios.py). 30-Jun-2026 (Q1FY2027 Reg 52(4)): 2.67% on the same "
         "excl.-basis - the series has improved for four consecutive years.", 0.0326,
         "held FLAT at the FY2025 level - the lending model forecasts the impairment "
         "ALLOWANCE stock, not the GNPA ratio, so no independent FY+1 GNPA forecast exists; "
         "using the latest audited figure is the conservative choice given the improving "
         "trend would otherwise be extrapolated without a documented basis"),
        (51, "NNPA", 0.0132,
         "FY2025 AR: net NPA 1.32% of loan assets (excl. interest accrued); 30-Jun-2026: "
         "1.23% on the same basis", 0.0132, "held flat, same reason as GNPA above"),
        (52, "RoA", 0.0310,
         "FY2025: return on average total assets = 3.10% (compute_ratios.py; reproduces the "
         "FY2025 AR's own 3.1% and CareEdge's independently published 3.11%, 30-Sep-2025)",
         0.0275,
         "Lending_Model P_and_L!C459.71 / Balance_Sheet!C average total assets 16,705.57 = "
         "2.75% FY2027E - a genuine projected decline versus FY2025, driven by the credit "
         "cost normalising off its FY2024-FY2026 write-back years (Assumptions!A_CC) and the "
         "opex ratio continuing its multi-year rise (Assumptions!A_OPEX), not by an "
         "aggressive input choice"),
        (53, "NIM", 0.0508,
         "FY2025: net interest margin on average total assets = 5.08% (compute_ratios.py; "
         "the FY2025 AR states 5.2% on an undisclosed basis - the gap is stated in "
         "compute_ratios.py's DISCLOSED_DIFFERENCES, not forced to reconcile)", 0.0549,
         "Lending_Model Credit_Ratios!C - FY2027E model NIM. CAVEAT carried from the model: "
         "finance costs are priced on OPENING debt (to avoid a circular reference through "
         "closing debt / retained earnings / PAT), which prices a growing book's financing "
         "at last year's lower stock and therefore mildly OVERSTATES this figure - the true "
         "FY+1 NIM is likely closer to flat on FY2025's 5.08% than to the model's 5.49%"),
        (54, "Cost-to-income", 0.2752,
         "FY2025: 27.52% (compute_ratios.py)", 0.2769,
         "Lending_Model FY2027E: opex 253.92 / (total income 1,950.47 less finance costs "
         "1,033.44) = 27.69% - effectively flat, consistent with the opex-ratio and "
         "yield/cost-of-funds assumptions both holding roughly steady"),
        (55, "Total assets / AUM", 14492, "FY2025 AR: reported AUM Rs 14,492 crore", 16637.83,
         "Lending_Model Book_Build!C - FY2027E closing AUM"),
    ]
    for r, _lab, latest, latest_note, expected, expected_note in rows:
        ws.cell(row=r, column=3, value=latest)
        ws.cell(row=r, column=4, value=expected)
        ws.cell(row=r, column=5, value="Latest: " + latest_note)
        ws.cell(row=r, column=6, value="Expected FY+1 (source: forecast model): "
                                        + expected_note)
        ws.column_dimensions["E"].width = 70
        ws.column_dimensions["F"].width = 70

    wb.save(OUT)
    _force_recalc(OUT)
    _cache_values(OUT)

    wbv = openpyxl.load_workbook(OUT, data_only=True)
    rat = wbv["Rating"]
    print("Saved:", OUT.name)
    print("  Sector type          :", rat["F6"].value)
    print("  Composite (financial+industry+business+mgmt, pre-notch):", rat["C6"].value)
    print("  Scale index pre-notch :", rat["C7"].value)
    print("  Support notches       :", rat["C8"].value)
    print("  Index after notches   :", rat["C9"].value)
    print("  Hard cap check        :", rat["C10"].value)
    print("  Index after caps      :", rat["C11"].value)
    print("  LONG-TERM ANCHOR GRADE :", rat["C12"].value)
    print("  SHORT-TERM ANCHOR GRADE:", rat["C13"].value)
    print("  Combined rating line   :", rat["C19"].value)
    sc = wbv["Scorecard"]
    print("\n  Financial-sector block (row 29-34), score column F:")
    for r in range(29, 35):
        print("   ", sc.cell(row=r, column=2).value, "->", sc.cell(row=r, column=6).value)


if __name__ == "__main__":
    main()
