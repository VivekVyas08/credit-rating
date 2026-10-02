"""Gate oracles for the Repco Home Finance Limited rating (GATES.md).

Each subcommand MEASURES a claim from the delivered artefacts and prints a success-only
token as its LAST line, exiting 0 only when every assertion passed. Any failure prints the
failing detail and exits 1, so a token can never appear on a failed run.

Usage: python verify_repco.py <gate> [--control]
       --control runs the securitisation oracle against a deliberately broken fixture and
       PASSES only if the oracle REJECTS it (positive control for G2).
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
RATIONALE = HERE / "Repco Home Finance Limited_24 Aug 2026_Rating Rationale.md"
SPREADS = HERE / "Repco Home Finance Limited_Financial_Spreads.xlsx"
MODEL = HERE / "Repco Home Finance Limited_Lending_Model.xlsx"
SCORECARD = HERE / "Rating_Scorecard_Repco Home Finance Limited.xlsx"
NBFC_METH = HERE.parents[1] / "Methodology" / "1003~NBFCs, Rating Methodology, Dec 2025.md"
HFC_METH = (HERE.parents[1] / "Methodology"
            / "651~Housing Finance Companies, Rating Methodology, Sep 2019 (Archived).md")
ICRA_REF = (HERE.parents[1] / "Report for Reference"
            / "Repco Home Finance Limited_25 Feb 2026_141196.md")

fails = []


def check(cond, msg):
    if not cond:
        fails.append(msg)
    return cond


def read(p):
    if not p.exists():
        fails.append("MISSING FILE %s" % p.name)
        return ""
    return p.read_text(encoding="utf-8", errors="replace")


# ------------------------------------------------------------------ G1 routing
def g_routing():
    meth = read(NBFC_METH)
    check(bool(meth), "NBFC methodology unreadable")
    check(re.search(r"housing finance compan(y|ies)\s*\(HFCs?\)", meth, re.I) is not None,
          "NBFC methodology's scope footnote does not name HFCs as covered")
    check(HFC_METH.exists(), "archived HFC methodology file missing from Methodology/")
    check("Archived" in HFC_METH.name, "HFC methodology file not marked Archived in its name")
    rat = read(RATIONALE)
    check("1003~NBFCs, Rating Methodology, Dec 2025" in rat or "NBFCs, Rating Methodology"
          in rat, "rationale does not name the NBFC methodology with vintage")
    check("Housing Finance Companies, Rating Methodology, Sep 2019" not in rat,
          "rationale cites the ARCHIVED HFC methodology - routing violation")
    return "ROUTING VERIFIED"


# ------------------------------------------------------------------ G2 / G2C securitisation
def g_securitisation(control=False):
    ref = read(ICRA_REF)
    check(bool(ref), "ICRA reference file for this issuer unreadable")
    check("SO)" in ref or "(SO)" in ref, "reference file does not carry a (SO) rating - "
          "expected a securitisation rationale")
    check("PRATHAM" in ref.upper(), "reference file does not name the PRATHAM trust")
    rat = read(RATIONALE)
    if control:
        # inject exactly the error this gate exists to catch: importing the SO rating as the
        # issuer's own corporate rating, with no explanatory text anywhere in the document
        rat = rat.replace(
            "**Repco Home Finance Limited: Ratings reaffirmed**",
            "**Repco Home Finance Limited: [ICRA]AAA(SO) reaffirmed**")
    aaa_so_as_issuer = re.search(
        r"^\*\*Repco Home Finance Limited:\s*\[ICRA\]AAA\(SO\)", rat, re.M)
    check(not aaa_so_as_issuer,
          "the (SO) securitisation rating appears to have been imported as the issuer's own "
          "headline rating")
    check("PRATHAM 12 2025" in rat, "rationale does not explain what the reference-library "
          "ICRA file actually is (the Pratham 12 2025 securitisation)")
    check("Originator" in rat or "originator" in rat,
          "rationale does not state RHFL's role (Originator/Servicer) in that transaction")
    check("[ICRA]AA- (Stable)" in rat or "[ICRA]AA-" in rat,
          "rationale does not carry an independent AA- corporate view distinct from the SO "
          "rating")
    return "SECURITISATION CORRECTLY EXCLUDED"


# ------------------------------------------------------------------ G3 handled by compute_ratios.py itself (CHECK: python compute_ratios.py)

# ------------------------------------------------------------------ G4 spreads
def g_spreads():
    if not SPREADS.exists():
        fails.append("MISSING WORKBOOK %s" % SPREADS.name)
        return "SPREADS OK"
    wb = openpyxl.load_workbook(SPREADS, data_only=True)
    for sh in ("Inputs", "RatioSuite", "TieOut", "RatingScore"):
        check(sh in wb.sheetnames, "spreads workbook missing sheet %s" % sh)
    if "TieOut" in wb.sheetnames:
        ws = wb["TieOut"]
        bad = tot = 0
        for r in range(6, ws.max_row + 1):
            c, d = ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value
            if isinstance(c, (int, float)) and isinstance(d, (int, float)):
                tot += 1
                if abs(c - d) > max(0.01, abs(d) * 0.005):
                    bad += 1
        check(tot >= 5, "only %d tie-out rows found" % tot)
        check(bad == 0, "%d of %d TieOut rows FAIL" % (bad, tot))
    return "SPREADS OK"


# ------------------------------------------------------------------ G5 Step 12
def g_step12():
    if not MODEL.exists():
        fails.append("MISSING MODEL %s" % MODEL.name)
        return "STEP 12 SATISFIED"
    wf = openpyxl.load_workbook(MODEL, data_only=False)
    wv = openpyxl.load_workbook(MODEL, data_only=True)
    need = ["Assumptions", "Book_Build", "P_and_L", "Balance_Sheet", "Credit_Ratios",
            "Debt_Breakup", "Stress_Case"]
    for sh in need:
        check(sh in wf.sheetnames, "lending model missing sheet %s" % sh)
    blanks = 0
    for sh in wf.sheetnames:
        for r in range(6, wf[sh].max_row + 1):
            for c in range(3, wf[sh].max_column + 1):
                fv = wf[sh].cell(row=r, column=c).value
                if isinstance(fv, str) and fv.startswith("="):
                    if wv[sh].cell(row=r, column=c).value is None:
                        blanks += 1
    check(blanks == 0, "%d formula cells in the lending model resolve to blank - a circular "
          "reference or broken link" % blanks)
    return "STEP 12 SATISFIED"


# ------------------------------------------------------------------ G6 Step 12b
def g_step12b():
    if not SCORECARD.exists():
        fails.append("MISSING SCORECARD %s" % SCORECARD.name)
        return "STEP 12B SATISFIED"
    wv = openpyxl.load_workbook(SCORECARD, data_only=True)
    check("Rating" in wv.sheetnames, "scorecard missing Rating sheet")
    grade = wv["Rating"]["C12"].value if "Rating" in wv.sheetnames else None
    check(isinstance(grade, str) and len(grade) > 0,
          "scorecard's Rating!C12 (long-term anchor grade) is blank")
    check(bool(re.match(r"^(AAA|AA[+-]?|A[+-]?|BBB[+-]?|BB[+-]?|B[+-]?|C)$", grade or ""),
          ), "scorecard anchor grade '%s' is not a recognised LT grade" % grade)
    return "STEP 12B SATISFIED"


# ------------------------------------------------------------------ G7 delivery gate
def g_gate():
    import subprocess
    r = subprocess.run(
        [sys.executable, str(HERE.parents[1] / "check_sections.py"), str(RATIONALE)],
        capture_output=True, text=True)
    out = r.stdout + r.stderr
    check(r.returncode == 0, "check_sections.py exited %d" % r.returncode)
    check("PASS" in out and "FAIL" not in out.split("PASS")[0],
          "check_sections.py output did not show a clean PASS: %s" % out[-800:])
    return "DELIVERY GATE CLEAN"


# ------------------------------------------------------------------ G8 rated amounts
def g_rated_amounts():
    rat = read(RATIONALE)
    check("to be confirmed by issuer" in rat.lower(),
          "no 'to be confirmed by issuer' marking found - rated amounts should be evidenced "
          "absent, not silently filled")
    check(not re.search(r"\*\*Total\*\*\s*\|\s*\*\*[\d,]+\.\d{2}\*\*", rat),
          "a numeric Total appears in the Summary of rating action table despite no rated "
          "amount being disclosed anywhere - likely an invented total")
    check("11,145.00" in rat and "CareEdge" in rat,
          "CareEdge's own rated amount is not shown as a labelled peer-agency reference")
    check("outstanding as on" not in rat.lower() and "outstanding as at" not in rat.lower()
          or "Annexure I" in rat,
          "an 'outstanding as on/at' figure appears where a rated limit is expected")
    return "RATED AMOUNTS SOUND"


# ------------------------------------------------------------------ G9 HFC metric set
def g_hfc_metrics():
    rat = read(RATIONALE)
    need = {
        "AUM": r"\bAUM\b",
        "gearing": r"\bgearing\b",
        "GNPA / Stage 3": r"\bGNPA\b|\bGross NPA\b|Stage 3",
        "CRAR": r"\bCRAR\b",
        "return measure": r"return on average (managed|total) assets|\bRoMA\b|\bRoAA\b",
    }
    missing = [k for k, p in need.items() if not re.search(p, rat, re.I)]
    check(not missing, "missing HFC metric set: %s" % ", ".join(missing))
    return "HFC METRIC SET PRESENT"


GATES = {"routing": g_routing, "securitisation": g_securitisation, "spreads": g_spreads,
         "step12": g_step12, "step12b": g_step12b, "gate": g_gate,
         "rated_amounts": g_rated_amounts, "hfc_metrics": g_hfc_metrics}

if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in GATES:
        print("usage: verify_repco.py {%s} [--control]" % "|".join(GATES))
        raise SystemExit(2)
    name = sys.argv[1]
    ctl = "--control" in sys.argv
    fn = GATES[name]
    try:
        token = fn(control=ctl) if fn.__code__.co_argcount else fn()
    except Exception as e:                                    # noqa: BLE001
        print("ORACLE ERROR: %s: %s" % (type(e).__name__, e))
        raise SystemExit(1)
    if ctl:
        if fails:
            print("CONTROL OK - oracle rejected the injected break (%d finding(s))"
                  % len(fails))
            raise SystemExit(0)
        print("CONTROL FAILED - oracle ACCEPTED a deliberately broken input; it proves nothing")
        raise SystemExit(1)
    for f in fails:
        print("FAIL:", f)
    if fails:
        print("%d assertion(s) failed" % len(fails))
        raise SystemExit(1)
    print(token)
