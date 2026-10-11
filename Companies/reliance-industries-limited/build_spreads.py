# -*- coding: utf-8 -*-
"""Generate RIL_Financial_Spreads.xlsx FROM compute_ratios.py (the script is the source of truth).
Tabs: Extraction / RatioSuite / TieOut / ExternalTieOut / Notes. Values are written LITERALLY (no
formulas) and render-verified by a data_only read-back."""
import importlib.util
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("cr", os.path.join(HERE, "compute_ratios.py"))
cr = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cr)

HDR, TITLE = Font(bold=True), Font(bold=True, size=12)
FILL = PatternFill("solid", fgColor="DDEBF7")
OK, BAD = PatternFill("solid", fgColor="E2F3E6"), PatternFill("solid", fgColor="FDECE0")
WRAP = Alignment(wrap_text=True, vertical="top")
YEARS = cr.YEARS
v = cr.v

SOURCES = {
    "PL": "Consolidated P&L (FY2023-24 AR for FY2023; FY2024-25 AR for FY2024; FY2025-26 AR p198)",
    "PL_PRINTED": "Consolidated P&L printed subtotals and attribution (same pages)",
    "OTHER_INCOME": "Consolidated note 26 Other income",
    "ONE_TIME_INCOME": "Q4 FY2026 media release note 5 (Rs 8,924 crore profit on sale of listed investments)",
    "FIN_COST": "Consolidated note 29 Finance costs",
    "CAPITALISED_INTEREST": "Note 29 footnote 'Net of Interest Capitalised'",
    "OTHER_EXP_SUB": "Consolidated note 30 Other expenses - subtotals (FY2025-26 AR)",
    "OTHER_EXP_MFG": "Note 30 manufacturing expenses lines (FY2023-FY2024 forex line from FY2023-24 AR)",
    "OTHER_EXP_SELL": "Note 30 selling and distribution lines",
    "OTHER_EXP_EST": "Note 30 establishment expenses lines",
    "NCA": "Consolidated balance sheet - non-current assets",
    "CA": "Consolidated balance sheet - current assets",
    "EQUITY": "Consolidated balance sheet - equity incl. NCI",
    "NCL": "Consolidated balance sheet - non-current liabilities",
    "CL": "Consolidated balance sheet - current liabilities",
    "BS_PRINTED": "Consolidated balance sheet printed totals",
    "DEBT_NOTES": "Notes 20 and 22 (current maturities, CP, current DPL, accrued interest)",
    "NC_BORROW_LINES": "Note 16 non-current borrowings by instrument (FY2025-26 AR)",
    "CB_LINES": "Note 20 current borrowings by instrument (FY2025-26 AR)",
    "DPL_AUCTIONS": "Note 17 footnote - DoT deferred payment liabilities by spectrum auction",
    "SUPPLIER_FINANCE": "Note 44 Supplier financing arrangement + cash-flow roll-forward",
    "CONTRACT_LIAB": "Note 23 Other current liabilities",
    "AR_NET_DEBT": "Note 35 Capital management - net gearing (each AR)",
    "AR_PUBLISHED": "Q4 FY2025 and Q4 FY2026 media releases (EBITDA, capex, ND/EBITDA)",
    "CF": "Consolidated statement of cash flows (each AR)",
    "CF_PRINTED": "Consolidated cash flow printed totals",
    "CF_OTHER": "Consolidated cash flow - remaining investing / financing lines",
    "DEBT_BRIDGE": "Change in liabilities arising from financing activities (below each cash flow)",
    "SEG_REV": "Audited consolidated segment information (Q4 FY2025 / Q4 FY2026 releases)",
    "SEG_EBITDA": "Segment results (EBITDA), same tables",
    "SEG_PRINTED": "Segment totals and reconciliation to PBT, same tables",
    "OPS": "Q4 FY2025 / Q4 FY2026 releases - Jio, Retail, O2C, Oil and Gas operating tables",
    "CONTINGENT": "Note 34 Contingent liabilities and commitments (FY2025-26 AR)",
    "PROMOTER_PCT": "Corporate governance report - distribution of shareholding",
}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Extraction"
ws["A1"] = "Reliance Industries Limited (consolidated) - extracted line items (Rs crore unless stated)"
ws["A1"].font = TITLE
ws["A2"] = ("Every value below is a literal copy of compute_ratios.py, which carries the source of each "
            "table. Nothing here is typed by hand. Blank = not disclosed / not extracted (see Notes).")
ws["A2"].alignment = WRAP
ws.row_dimensions[2].height = 30
ws.append([])
ws.append(["Table", "Line item"] + YEARS + ["Source"])
for c in ws[4]:
    c.font, c.fill = HDR, FILL
for name in SOURCES:
    for key, series in getattr(cr, name).items():
        ws.append([name, key] + [None if x is None else float(x) for x in series] + [SOURCES[name]])
for col, w in zip("ABCDEFG", [22, 50, 13, 13, 13, 13, 70]):
    ws.column_dimensions[col].width = w

R = cr.compute()
ws2 = wb.create_sheet("RatioSuite")
ws2["A1"] = "CRISIL-convention ratio suite - consolidated (full consolidation)"
ws2["A1"].font = TITLE
ws2["A2"] = ("PBDIT includes recurring other income and excludes the FY2026 one-time gain (8,924); interest "
             "includes capitalised interest; adjusted debt = borrowings + leases + DoT deferred payment "
             "liabilities (+ supplier finance in the 'incl. SF' rows); adjusted networth incl. NCI.")
ws2["A2"].alignment = WRAP
ws2.row_dimensions[2].height = 45
ws2.append([])
ws2.append(["Metric"] + YEARS)
for c in ws2[4]:
    c.font, c.fill = HDR, FILL
for key in R[YEARS[-1]]:
    ws2.append([key] + [None if R[y].get(key) is None else round(float(R[y][key]), 2) for y in YEARS])
ws2.column_dimensions["A"].width = 66
for col in "BCDE":
    ws2.column_dimensions[col].width = 14

T = cr.tieouts(R)
ws3 = wb.create_sheet("TieOut")
ws3["A1"] = "Internal tie-outs - every check compares a computation with a figure a document prints"
ws3["A1"].font = TITLE
ws3.append([])
ws3.append(["Check", "Computed", "Disclosed", "Difference", "Tolerance", "Result", "Source"])
for c in ws3[3]:
    c.font, c.fill = HDR, FILL
npass = 0
for d, comp, disc, tol, src in T:
    ok = abs(comp - disc) <= tol
    npass += ok
    ws3.append([d, round(comp, 3), round(disc, 3), round(comp - disc, 3), tol, "PASS" if ok else "FAIL", src])
    ws3.cell(row=ws3.max_row, column=6).fill = OK if ok else BAD
ws3.append([])
ws3.append(["TOTAL", None, None, None, None, "%d/%d PASS" % (npass, len(T))])
ws3.append(["Mutation coverage (tools/tieout_mutation.py): 640 of 675 inputs change a tie-out = 94.8%; "
            "the 35 that do not are named in Notes / DISCLOSED_DIFFERENCES."])
for col, w in zip("ABCDEFG", [74, 13, 13, 11, 10, 9, 60]):
    ws3.column_dimensions[col].width = w

# ---- External tie-out vs screener.in consolidated view (read 10-Oct-2026)
ws4 = wb.create_sheet("ExternalTieOut")
ws4["A1"] = "External tie-out vs screener.in consolidated (cross-check ONLY; read 10-Oct-2026)"
ws4["A1"].font = TITLE
ws4.append([])
ws4.append(["Item", "Year", "screener.in", "Engine (rebuilt on screener's basis)", "Difference", "Basis note"])
for c in ws4[3]:
    c.font, c.fill = HDR, FILL


def loss_ppe(p):
    return v(cr.OTHER_EXP_EST, "Loss on sale/discard of PPE", p)


def sales_tax(p):
    return v(cr.OTHER_EXP_SELL, "Sales tax / VAT", p)


SCR = []
for p, s_sales in (("FY2025", 962820), ("FY2026", 1055780)):
    SCR.append(("Sales", p, s_sales, v(cr.PL, "Revenue from operations", p) - v(cr.PL, "Excise duty", p) - sales_tax(p),
                "screener nets excise duty and sales tax/VAT (note 30) from revenue"))
for p, s_op in zip(YEARS, (142318, 162498, 165598, 179065)):
    SCR.append(("Operating profit", p, s_op, R[p]["OPBDIT excl. all other income"] + loss_ppe(p),
                "screener moves loss on sale/discard of PPE into other income"))
for p, s_oi in zip(YEARS, (12020, 15792, 17824, 28846)):
    extra = (v(cr.PL, "Share of associates and JVs", p) + v(cr.PL, "Profit from discontinued operations", p)
             if p == "FY2023" else 0)
    SCR.append(("Other income", p, s_oi, v(cr.PL, "Other income", p) - loss_ppe(p) + extra,
                "less loss on PPE; FY2023 also carries associates (24) + discontinued profit (418)"))
for p, s_i in zip(YEARS, (19571, 23118, 24269, 27061)):
    SCR.append(("Interest", p, s_i, v(cr.PL, "Finance costs", p), "P&L finance costs (CRISIL adds capitalised)"))
for p, s_d in zip(YEARS, (40303, 50832, 53136, 57688)):
    SCR.append(("Depreciation", p, s_d, v(cr.PL, "Depreciation and amortisation", p), "P&L"))
for p, s_n in zip(YEARS, (74088, 79020, 81309, 95754)):
    SCR.append(("Net profit", p, s_n, R[p]["PAT (incl. associates)"], "profit for the year"))
for p, s_b in zip(YEARS, (451664, 350719, 374313, 402962)):
    eng = (cr.adj_debt(p) if p == "FY2023" else
           cr.borrowings(p) + cr.leases(p) + v(cr.DEBT_NOTES, "Current maturities of deferred payment liabilities", p))
    SCR.append(("Borrowings", p, s_b, eng,
                "FY2023: borrowings + leases + ALL DoT liabilities; FY2024-26: borrowings + leases + current "
                "DoT instalment only (non-current DoT sits in screener's 'other liabilities')"))
for p, s_r in zip(YEARS, (709106, 786715, 829668, 890498)):
    SCR.append(("Reserves", p, s_r, v(cr.EQUITY, "Other equity", p), "other equity"))
for p, s_c in zip(YEARS, (115032, 158788, 178703, 192113)):
    SCR.append(("Cash from operating activity", p, s_c, R[p]["CFO"], "cash flow"))
for p, s_t in zip(YEARS, (1605882, 1755048, 1949713, 2177546)):
    SCR.append(("Total assets", p, s_t, v(cr.BS_PRINTED, "Total assets", p) - v(cr.NCA, "Deferred tax assets", p),
                "screener nets deferred tax assets against liabilities"))
next_ok = 0
for item, y, s, e, note in SCR:
    diff = e - s
    ok = abs(diff) < 1
    next_ok += ok
    ws4.append([item, y, s, round(e, 2), round(diff, 2), ("PASS - " if ok else "CHECK - ") + note])
    ws4.cell(row=ws4.max_row, column=6).fill = OK if ok else BAD
ws4.append([])
ws4.append(["TOTAL", None, None, None, None, "%d/%d reconcile exactly" % (next_ok, len(SCR))])
for col, w in zip("ABCDEF", [30, 9, 13, 18, 11, 100]):
    ws4.column_dimensions[col].width = w

ws5 = wb.create_sheet("Notes")
ws5["A1"] = "Basis of preparation and disclosed differences"
ws5["A1"].font = TITLE
ws5.append([])
for title, text in cr.DISCLOSED_DIFFERENCES:
    ws5.append([title])
    ws5.cell(row=ws5.max_row, column=1).font = HDR
    ws5.append([text])
    ws5.cell(row=ws5.max_row, column=1).alignment = WRAP
    ws5.row_dimensions[ws5.max_row].height = 80
    ws5.append([])
ws5.column_dimensions["A"].width = 130

out = os.path.join(HERE, "RIL_Financial_Spreads.xlsx")
wb.save(out)
wb2 = openpyxl.load_workbook(out, data_only=True)
cells = formulas = 0
for s in wb2.sheetnames:
    for row in wb2[s].iter_rows():
        for c in row:
            if c.value is not None:
                cells += 1
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
print("saved %s; sheets=%s; non-empty cells=%d; formula cells=%d; tie-outs %d/%d PASS; external %d/%d"
      % (out, wb2.sheetnames, cells, formulas, npass, len(T), next_ok, len(SCR)))
for item, y, s, e, note in SCR:
    if abs(e - s) >= 1:
        print("  CHECK", item, y, s, round(e, 2))
