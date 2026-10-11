"""Fills Rating_Scorecard_Reliance Industries Limited.xlsx (Step 12b; Stage 2 = Step 17c).

Closest scorecard sector: 'Generic corporate' - RIL is a diversified conglomerate (FY2026 segment
EBITDA: Digital 39%, O2C 31%, Retail 14%, Oil & Gas 10%, Others 6%); no single scorecard sector fits.
Historical = audited consolidated FY2024-FY2026 from compute_ratios.py on CRISIL definitions.
FY+1/FY+2 = COPIED (never linked) from RIL_Financial_Model_10Oct2026.xlsx base case."""
import importlib.util
import shutil
import sys
from pathlib import Path

import openpyxl

CASE = Path(__file__).resolve().parent
ROOT = CASE.parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from make_spreads import _cache_values, _force_recalc  # noqa: E402

spec = importlib.util.spec_from_file_location("cr", CASE / "compute_ratios.py")
E = importlib.util.module_from_spec(spec)
spec.loader.exec_module(E)
R = E.compute()
v, tot = E.v, E.tot
OUT = CASE / "Rating_Scorecard_Reliance Industries Limited.xlsx"
MODEL = CASE / "RIL_Financial_Model_10Oct2026.xlsx"
Y3 = ["FY2024", "FY2025", "FY2026"]


def model_row(tab, label):
    ws = openpyxl.load_workbook(MODEL, data_only=True)[tab]
    for row in ws.iter_rows(values_only=True):
        if row[0] == label:
            return row[1:10]  # FY2023..FY2031E
    raise KeyError(label)


rev = model_row("Income Statement", "Revenue from operations (net of GST)")
pbdit = model_row("Income Statement", "PBDIT (CRISIL: incl. recurring other income, excl. one-time)")
pat = model_row("Income Statement", "PROFIT FOR THE YEAR")
capex = model_row("Cash Flow", "Capex - PPE, spectrum, intangibles")
div = model_row("Cash Flow", "Dividends paid (owners + NCI)")
adj = model_row("Credit Ratios", "Adjusted total debt (borrowings + leases + DoT liabilities)")
# index 3 = FY2026, 4 = FY2027E, 5 = FY2028E

shutil.copy(ROOT / "Rating_Scorecard.xlsx", OUT)
wb = openpyxl.load_workbook(OUT)
ws = wb["Inputs"]
ws["C4"], ws["C5"], ws["C6"], ws["C7"] = "Reliance Industries Limited", "CRISIL", "Generic corporate", "10-Oct-2026"

hist = [
    (12, [v(E.PL, "Revenue from operations", p) for p in Y3], "Revenue from operations (net of GST), consolidated P&L"),
    (13, [R[p]["PBDIT (CRISIL)"] for p in Y3],
     "CRISIL PBDIT: PBT + finance costs + D&A, i.e. INCLUDING recurring other income and EXCLUDING the "
     "FY2026 one-time Rs 8,924 crore profit on sale of listed investments; company EBITDA 1,78,290 / "
     "1,83,422 / 2,07,911 (engine tie-outs)"),
    (14, [v(E.PL, "Depreciation and amortisation", p) for p in Y3], "consolidated P&L"),
    (15, [R[p]["Interest (incl. capitalised)"] for p in Y3],
     "CRISIL interest = P&L finance costs (incl. lease interest and forex loss) + interest capitalised "
     "15,222 / 18,283 / 16,091 (note 29 footnote)"),
    (16, [R[p]["PAT (incl. associates)"] for p in Y3], "profit for the year incl. NCI share and associates"),
    (17, [R[p]["Adjusted total debt"] for p in Y3],
     "CRISIL adjusted debt = borrowings + leases + DoT deferred spectrum liabilities. Supplier finance "
     "(note 44: 1,03,716 FY2025 / 76,377 FY2026) is NOT included here; including it, FY2026 gearing is "
     "0.53x and debt/PBDIT 2.91x (engine)"),
    (18, [R[p]["Cash and marketable securities"] for p in Y3],
     "cash + current investments + other marketable securities (the AR's own note-35 definition); no "
     "earmarked balances identified at this level of disclosure"),
    (19, [R[p]["Adjusted networth (incl. NCI)"] for p in Y3],
     "adjusted networth = total equity incl. non-controlling interest; DTL excluded; no revaluation "
     "reserve; goodwill/spectrum not deducted (acquired operating licences - CRISIL assesses case-specifically)"),
    (20, [E.total_liabilities(p) for p in Y3], "total liabilities, consolidated balance sheet"),
    (21, [tot(E.CA, p) for p in Y3], "total current assets"),
    (22, [tot(E.CL, p) for p in Y3], "total current liabilities"),
    (23, [v(E.CF, "Capex (PPE, spectrum, intangibles)", p) for p in Y3], "consolidated cash flow"),
    (24, [v(E.CF, "Dividend paid to owners", p) + v(E.CF, "Dividend paid to NCI", p) for p in Y3],
     "consolidated cash flow (owners + NCI)"),
]
for r, vals, note in hist:
    for j, x in enumerate(vals):
        ws.cell(r, 3 + j, round(x, 2))
    ws.cell(r, 6, note)

proj = [
    (29, rev[4] / rev[3] - 1, rev[5] / rev[4] - 1, "Revenue Model (segment growth: O2C 5%/3%, Retail 12%, Digital 13%/12%)"),
    (30, pbdit[4] / rev[4], pbdit[5] / rev[5], "Income Statement (segment EBITDA + treasury yield + unallocable)"),
    (31, capex[4], capex[5], "Assumptions (cash capex 1,50,000 / 1,45,000)"),
    (32, adj[4] - adj[3], adj[5] - adj[4],
     "Credit Ratios (no funding-gap draw in the base case; only DoT instalments 4,962 / 5,300 reduce debt)"),
    (33, 0.08, 0.08, "Assumptions (gross 8.0% on borrowings + DoT liabilities, part capitalised)"),
    (34, 0.23, 0.23, "Assumptions (tax 23%)"),
    (35, div[4] / pat[4], div[5] / pat[5],
     "Cash Flow / Income Statement: FY+1 = Rs 6/share for FY2026 (8,119, note 45) + NCI ~450; later ~10% of PAT"),
]
for r, a, b, src in proj:
    ws.cell(r, 3, round(a, 4))
    ws.cell(r, 4, round(b, 4))
    ws.cell(r, 5, "source: forecast model (RIL_Financial_Model_10Oct2026.xlsx, %s)" % src)

ws["C40"] = 1
ws["F40"] = ("1 - leadership in every core business: O2C throughput 80.0 MMT (FY2026) at an integrated "
             "refining-petrochemical complex; Jio 524.4 million subscribers (Q4 FY2026) and Q4 ARPU Rs 214.0; "
             "Retail 20,160 stores, 78.3 mn sq ft, 387 mn registered customers; four segments, none above "
             "39.5% of FY2026 segment EBITDA (engine) - the diversity CRISIL treats as a business-risk strength.")
ws["C41"] = 2
ws["F41"] = ("2 - promoter group 49.11% (FY2026 AR); unmodified audit opinions with KAMs on litigation, a "
             "Level-3 valuation (Rs 77,852 crore JDFPL preference shares) and Jio revenue/IT; risk appetite "
             "is high - capex 1.31-1.44 lakh crore a year and new-energy projects large relative to most "
             "Indian corporates (CRISIL: projects larger than operations viewed unfavourably), but funded "
             "within a stated investment-grade capital framework (note 35). Not 1: the SEBI RPL disgorgement "
             "(Rs 447.27 crore + 12% interest) awaits the Supreme Court; KG D6 USD 1.87 billion demand open.")
ws["C42"] = 0
ws["F42"] = "0 - RIL is the group flagship; no parent support claimed and no downward linkage (no weaker parent)."
ws["C43"] = "Superior"
ws["F43"] = ("Step 11 arithmetic: cash and marketable securities Rs 2,49,704 crore (31-Mar-2026) against debt "
             "falling due within a year (current maturities 25,965 + DoT 4,962 + leases 7,381 = 38,308) "
             "plus CP 13,186 - cover 4.85x; Q1 FY2027 cash Rs 2,46,791 crore.")
ws["C44"] = "Stable"
# ---- Stage 2: Rating Committee decision, Step 17c (Committee_Minutes_10Oct2026.md) ----
ws["C61"] = 1
ws["D61"] = ("+1 Business / franchise strength beyond the financial composite (Committee, 10-Oct-2026): the "
             "composite scores a conglomerate as one 'Generic corporate' and cannot see diversity - no "
             "segment above 39.5% of FY2026 segment EBITDA, consumer share 46.7% -> 53.4% (FY2024-26), and "
             "the FY2025 O2C fall (-7,401) more than offset by Digital and Retail (+10,312). Liquidity and "
             "flexibility are beyond the template's bands: cash and marketable securities Rs 2,49,704 crore "
             "= 4.85x debt due within a year plus CP; market cap ~Rs 15.8 lakh crore = 3.2x adjusted debt. "
             "Base case: adjusted debt/PBDIT 2.30x FY2027 -> 1.60x FY2031 with no new borrowing. Known "
             "template bias: under-rates AA+/AAA (back-test 18-Jul-2026). Not applied as a hard-cap "
             "disapplication - no cap binds.")
wb.save(OUT)
_force_recalc(OUT)
_cache_values(OUT)
wbv = openpyxl.load_workbook(OUT, data_only=True)
rat, sc = wbv["Rating"], wbv["Scorecard"]
for r in (10, 11, 13, 15, 16, 17, 32, 34, 38):
    print("  Rating!C%-3d %-55s %s" % (r, str(rat.cell(r, 2).value)[:55], rat.cell(r, 3).value))
for r in range(5, 45):
    a, b, c = sc.cell(r, 2).value, sc.cell(r, 3).value, sc.cell(r, 5).value
    if a:
        print("  Scorecard!%-3d %-55s %s %s" % (r, str(a)[:55], b, c))
