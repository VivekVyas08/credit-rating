"""Step 12b - fill the native Rating_Scorecard template for Alembic Limited.

Historicals are read from compute_ratios.py (never retyped). Projections are COPIED from
Alembic_Limited_Sensitivity_Model_31Aug2026.md, which is the scaled Step 12 artefact recorded
in case.json's step12_scaling_decision.
"""

import shutil
import sys

from openpyxl import load_workbook

import compute_ratios as C

SRC = "../../Rating_Scorecard.xlsx"
DST = "Rating_Scorecard_Alembic_Limited.xlsx"
HIST = ["FY2024", "FY2025", "FY2026"]

# Base case, from the focused sensitivity. Each driver carries its derivation.
PROJ = {
    # Real estate broadly flat (185.35 in FY2026, 182.14 in FY2025) with API compounding off a
    # small base (+44.6% in FY2026). A blended 8% is below the two-year actual of +23.5% CAGR
    # because API cannot sustain 45% and real estate is lumpy across project completions.
    "revenue_growth":  (0.08, 0.08),
    # OPBDIT/OI was 36.95 -> 41.78 -> 37.81. Held at the FY2026 level rather than the FY2025
    # peak; the segment note attributes the easing to real-estate project mix.
    "ebitda_margin":   (0.375, 0.375),
    # FY2026 actual capex Rs 76.65 crore; capital commitments outstanding Rs 13.52 crore.
    "capex":           (70.0, 70.0),
    # DERIVED FROM THE FUNDING GAP, not assumed flat (failure mode 18):
    #   internal accrual  = PAT ~125 + depreciation ~13   = ~138
    #   less dividend     = ~61 (49% payout, FY2026 actual Rs 61.41 crore)
    #   less capex        = 70
    #   => ~+7 crore surplus, so no new external debt is needed in either year and the only
    #      movement is scheduled repayment. FY2026 actual repayment was Rs 3.81 crore.
    "net_new_debt":    (-4.0, -4.0),
    # Implied rate on the average book: finance cost Rs 1.57 crore / average debt Rs 12.55
    # crore. High because the book is tiny and the line includes non-borrowing bank charges.
    "interest_rate":   (0.125, 0.125),
    # FY2026 implied rate 15.2%. The Section 80M shelter is a fixed rupee amount, so it covers
    # a shrinking share of a growing profit; the rate is nudged up rather than held.
    "tax_rate":        (0.16, 0.16),
    # FY2026 actual: Rs 61.41 crore paid on Rs 124.54 crore PAT.
    "dividend_payout": (0.49, 0.49),
}

QUALITATIVE = {
    "business": (4, "Real estate is 80% of revenue (Rs 185.35 crore of Rs 231.93 crore) and 93% "
                    "of segment profit, and it is a single-city landholder-developer business "
                    "in Vadodara - geographically concentrated and lumpy across project "
                    "completions. API adds Rs 46.59 crore, growing 44.6%, but off a small base. "
                    "Operating scale is the binding constraint on this credit and is scored as "
                    "such; margins (37.81% OPBDIT) are strong but do not offset modest size."),
    "management": (3, "FY2026 audit opinion unmodified on both bases, with an Emphasis of "
                      "Matter on the electricity-duty litigation and a Key Audit Matter on "
                      "litigations and provisions. Not a 2, for two documented reasons: FY2023 "
                      "drew a QUALIFIED opinion on a Rs 10.61 crore transfer from General "
                      "Reserve to the P&L, and related-party dealing is dense - Rs 38.67 crore "
                      "of rent income from the wholly-owned subsidiary and Rs 38.86 crore of "
                      "dividend to the 63.06% holding company in FY2026."),
    "support": (0, "No notch in either direction. Nirayu Private Limited holds 63.06% but NO "
                   "financial information on it exists in the document set - searched - so its "
                   "credit profile cannot be assessed and no uplift may be taken. Equally no "
                   "cap: AL needs no support in any evidenced scenario (net cash Rs 81.38 "
                   "crore, quoted investments 365x total debt)."),
    "liquidity": ("Superior",
                  "On-balance-sheet liquidity of Rs 92.02 crore (cash Rs 2.97 crore + current "
                  "investments Rs 89.05 crore) against total debt of Rs 10.64 crore - 8.6x "
                  "cover. Recurring dividend income of Rs 62.95 crore alone covers total debt "
                  "5.9x and finance cost 40x. Behind that sit Rs 3,885.89 crore of listed, "
                  "saleable quoted investments. No encumbrance on the cash or current "
                  "investments is disclosed in note 12, 13 or 10 - checked, not assumed "
                  "(failure mode 34). The arithmetic does not admit a weaker descriptor."),
    "outlook": ("Stable", "See the rationale."),
}


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                                     # noqa: BLE001
        pass

    shutil.copyfile(SRC, DST)
    wb = load_workbook(DST)
    ws = wb["Inputs"]
    R = C.compute()
    I = C.IDX if hasattr(C, "IDX") else {y: i for i, y in enumerate(C.YEARS)}

    ws["C4"] = "Alembic Limited"
    ws["C5"] = "ICRA"
    ws["C6"] = "Real Estate / Construction"
    ws["C7"] = "31-Aug-2026"

    rows = {
        12: lambda y: R[y]["Operating income"],
        13: lambda y: R[y]["OPBDIT"],
        14: lambda y: C.SA["depreciation"][I[y]],
        15: lambda y: C.SA["finance_cost"][I[y]],
        16: lambda y: R[y]["PAT"],
        17: lambda y: R[y]["Total debt"],
        18: lambda y: C.SA["cash"][I[y]] + C.SA["curr_inv"][I[y]],
        19: lambda y: R[y]["Tangible net worth"],
        20: lambda y: C.SA["total_assets"][I[y]] - R[y]["Tangible net worth"],
        21: lambda y: R[y]["Current assets"],
        22: lambda y: R[y]["Current liabilities"],
        23: lambda y: R[y]["Capex"],
        24: lambda y: R[y]["Dividend paid"],
    }
    for r, fn in rows.items():
        for j, y in enumerate(HIST):
            ws.cell(row=r, column=3 + j, value=round(fn(y), 2))

    for r, key in [(29, "revenue_growth"), (30, "ebitda_margin"), (31, "capex"),
                   (32, "net_new_debt"), (33, "interest_rate"), (34, "tax_rate"),
                   (35, "dividend_payout")]:
        for j in range(2):
            ws.cell(row=r, column=3 + j, value=PROJ[key][j])

    ws["C39"] = None                       # no industry override; sector base = 4.5
    ws["C40"] = QUALITATIVE["business"][0]
    ws["C41"] = QUALITATIVE["management"][0]
    ws["C42"] = QUALITATIVE["support"][0]
    ws["C43"] = QUALITATIVE["liquidity"][0]
    ws["C44"] = QUALITATIVE["outlook"][0]

    wb.save(DST)
    print("wrote %s" % DST)

    try:
        sys.path.insert(0, "../..")
        import _xlread
        cells = _xlread.load(DST)
        print("Rating!C12 anchor grade: %s" % cells.get(("RATING", "C", 12)))
        for rr in range(4, 20):
            lab = cells.get(("RATING", "B", rr))
            if lab:
                print("   %-38s %s" % (lab, cells.get(("RATING", "C", rr))))
    except Exception as e:                                                # noqa: BLE001
        print("anchor read-back unavailable (%s)" % e)

    print("\nQualitative scores and their justification:")
    for k in ("business", "management", "support", "liquidity"):
        v, why = QUALITATIVE[k]
        print("  %-11s %-9s %s" % (k, v, why))


if __name__ == "__main__":
    main()
