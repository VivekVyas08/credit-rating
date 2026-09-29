"""Rebuild Rating_Scorecard.xlsx's Rating sheet as a THREE-STAGE sheet whose HEADLINE is the
ASSIGNED rating - the same rating the rationale states - with the financial anchor and the
bridge between them shown underneath.

WHY THIS EXISTS
---------------
The sheet used to output the financial anchor and nothing else, formatted exactly like an
assigned rating. Delivered beside a rationale it produced two different grades on two documents
with no visible connection - Vodafone Idea (anchor B-, rating A-) and Godrej Properties
(anchor B+, rating AA+). Operator instruction, 02-Sep-2026: "make it in a way that they both
are the same and they are both linked to each other".

The link is not cosmetic. It is a computed chain:

    STAGE 1  financial anchor        (audited financials only - the screening view)
       |     + committee overlay      (Stage 2: named, justified notch adjustments)
       v
    STAGE 3  ASSIGNED RATING          (the headline, and what the rationale must state)

Nothing analytical is given up. The anchor is still computed standalone and still visible; the
divergence is still forced into the open. What changes is that the divergence is now COMPUTED
on the sheet rather than typed into prose - which matters, because the first hand-typed notch
count in this project was wrong (the Godrej Properties rationale said 9 notches when the
chain gives 12; the Vodafone Idea report said 12 when it was 9).

TWO SUBSTANTIVE FIXES SHIP WITH THE RESTRUCTURE
-----------------------------------------------
1. HARD CAPS BECOME DISAPPLICABLE, WITH A WRITTEN REASON. The caps test accrual metrics -
   "interest coverage < 1x" forces the grade to B+. For a real-estate developer under Ind AS
   project-completion accounting, accrual coverage is structurally negative in most years while
   cash coverage is ~57x; the cap fires on an artefact and moved Godrej Properties from A+ to
   B+, manufacturing 9 of its 12 notches of divergence. The cap is not deleted - it is now
   disapplied only by setting Inputs!C67 to "N" WITH a written reason in Inputs!D67, which the
   sheet prints and the delivery gate reads.
2. THE OVERLAY IS AN INPUT WITH A MANDATORY JUSTIFICATION. Each Stage 2 adjustment carries its
   own reason cell. A notch with no reason is a defect the gate reports.

    python tools/restructure_scorecard.py            # rebuild the template
    python tools/restructure_scorecard.py --check    # print the new layout, change nothing
"""

import argparse
import shutil
import sys
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "Rating_Scorecard.xlsx"

NAVY = "1B2A4A"
TEAL = "0E7C7B"
AMBER = "FFF4E5"
GREY = "F2F4F7"
WHITE = "FFFFFF"

# ---------------------------------------------------------------- Inputs sheet overlay block
OVERLAY_ROWS = [
    (60, "Sector-metric applicability",
     "Where a scorecard metric is structurally distorted for this sector (e.g. accrual "
     "interest coverage for a project-completion real-estate developer), state it here."),
    (61, "Business / franchise strength beyond the financial composite",
     "Market position, scale or diversification the financial ratios cannot see."),
    (62, "Parent / group support beyond the Step 12b notch",
     "Only the INCREMENT over Inputs!C42, so support is never counted twice."),
    (63, "Peer benchmarking / calibration",
     "Where the anchor sits against this agency's published grades for close peers."),
    (64, "Other (name it in the justification)",
     "Anything else the Committee relied on. Unnamed adjustments are a defect."),
]


def build_inputs_block(ws):
    """Committee-overlay inputs at rows 58-67 (previously empty)."""
    bold = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor=NAVY)
    sub = Font(bold=True, color=NAVY)
    amber = PatternFill("solid", fgColor=AMBER)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["B58"] = ("COMMITTEE OVERLAY (Step 17c) - the bridge from the financial anchor to the "
                 "ASSIGNED rating")
    ws["B58"].font = bold
    ws["B58"].fill = hdr_fill

    ws["B59"] = "Adjustment"
    ws["C59"] = "Notches (+ improves)"
    ws["D59"] = "Justification - REQUIRED whenever notches <> 0"
    for c in ("B59", "C59", "D59"):
        ws[c].font = sub

    for row, label, hint in OVERLAY_ROWS:
        ws.cell(row=row, column=2, value=label)
        ws.cell(row=row, column=3, value=0).fill = amber
        ws.cell(row=row, column=4, value="").fill = amber
        ws.cell(row=row, column=6, value=hint).alignment = wrap

    ws["B66"] = "Committee direct override - assigned index (blank = anchor + overlay)"
    ws["C66"] = ""
    ws["C66"].fill = amber
    ws["D66"] = ""
    ws["D66"].fill = amber
    ws["F66"] = ("Use ONLY where the Committee sets the grade directly rather than by notching "
                 "from the anchor. 0 = AAA ... 16 = C. A reason in D66 is required.")
    ws["F66"].alignment = wrap

    ws["B67"] = "Hard caps apply? (Y / N)"
    ws["C67"] = "Y"
    ws["C67"].fill = amber
    ws["D67"] = ""
    ws["D67"].fill = amber
    ws["F67"] = ("'N' ONLY where a capped metric is structurally meaningless for this sector - "
                 "e.g. accrual interest coverage for a project-completion real-estate "
                 "developer, whose cash coverage is strong while its accrual coverage is "
                 "negative by construction. A written reason in D67 is REQUIRED and the "
                 "delivery gate reads it.")
    ws["F67"].alignment = wrap


# ---------------------------------------------------------------------- Rating sheet rebuild
def build_rating_sheet(ws):
    """Wipe and rewrite the Rating sheet in the three-stage layout."""
    # The original sheet merges its banner cells; a MergedCell is read-only, so every merge
    # must be released before the sheet can be rewritten.
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    for row in ws["A1:H60"]:
        for cell in row:
            cell.value = None

    hdr = Font(bold=True, color="FFFFFF", size=11)
    navy_fill = PatternFill("solid", fgColor=NAVY)
    teal_fill = PatternFill("solid", fgColor=TEAL)
    big = Font(bold=True, size=16, color=NAVY)
    sub = Font(bold=True, color=NAVY)
    mut = Font(italic=True, size=9, color="6B7891")
    wrap = Alignment(wrap_text=True, vertical="top")
    grey = PatternFill("solid", fgColor=GREY)

    # ---------------------------------------------------------------- headline (Stage 3 echo)
    ws["B2"] = "ASSIGNED RATING   -   this is the rating the rationale must state"
    ws["B2"].font = hdr
    ws["B2"].fill = teal_fill

    ws["B3"] = "Long-term"
    ws["C3"] = "=C36"
    ws["C3"].font = big
    ws["B4"] = "Short-term"
    ws["C4"] = "=C37"
    ws["C4"].font = big
    ws["B5"] = "Combined"
    ws["C5"] = "=C38"
    ws["C5"].font = Font(bold=True, size=12, color=NAVY)

    ws["B6"] = ('="Issuer: "&Inputs!C4&"   |   Agency: "&Inputs!C5&"   |   Sector: "&Inputs!C6')
    ws["B6"].font = sub
    ws["E6"] = "Sector type"
    ws["F6"] = ("=INDEX(Ref_Weights!$H$5:$H$26,MATCH(Inputs!$C$6,Ref_Weights!$A$5:$A$26,0))")

    ws["B7"] = ("Derived below: STAGE 1 financial anchor  ->  STAGE 2 committee overlay  ->  "
                "STAGE 3 assigned rating. Every step is computed on this sheet.")
    ws["B7"].font = mut
    ws["B7"].alignment = wrap

    # ------------------------------------------------------------------------ STAGE 1: anchor
    ws["B9"] = "STAGE 1 - FINANCIAL ANCHOR   (screening view, computed from audited financials alone)"
    ws["B9"].font = hdr
    ws["B9"].fill = navy_fill

    ws["B10"] = "Adjusted composite score"
    ws["C10"] = "=Scorecard!E24"
    ws["B11"] = "Scale index before notches (0 = AAA ... 16 = C)"
    ws["C11"] = "=MIN(16,MAX(0,ROUND(16*((C10-1)/5)^1.4,0)))"
    ws["D11"] = ("back-tested map (18-Jul-2026): power p=1.4, zero bias, reliable to about "
                 "+/-2 notches as a SCREEN")
    ws["D11"].font = mut
    ws["B12"] = "Support notches entered at Step 12b (+ improves)"
    ws["C12"] = "=Inputs!C42"
    ws["B13"] = "Index after support notches"
    ws["C13"] = "=MIN(16,MAX(0,C11-C12))"

    ws["B14"] = "Hard cap tested"
    ws["C14"] = ('=IF(Inputs!C67="N","DISAPPLIED for this issuer - reason in Inputs!D67",'
                 'IF(F6="Corporate",'
                 'IF(Inputs!E19<0,"Negative net worth: capped at B-",'
                 'IF(Ratios!E14<1,"Coverage <1x: capped at B+","None")),'
                 'IF(Inputs!C49<0.09,"CRAR below 9%: capped at BB+","None")))')
    # "" (empty string) is indistinguishable from "never recalculated" to openpyxl's
    # data_only read - the same C30 trap. Render a placeholder instead.
    ws["D14"] = '=IF(Inputs!C67="N",Inputs!D67,"(cap applies - not disapplied)")'
    ws["D14"].font = mut
    ws["D14"].alignment = wrap

    ws["B15"] = "ANCHOR INDEX (after caps)"
    ws["C15"] = ('=IF(Inputs!C67="N",C13,'
                 'IF(F6="Corporate",'
                 'IF(Inputs!E19<0,MAX(C13,15),'
                 'IF(Ratios!E14<1,MAX(C13,13),C13)),'
                 'IF(Inputs!C49<0.09,MAX(C13,10),C13)))')
    ws["C15"].font = sub
    ws["B16"] = "Anchor long-term grade"
    ws["C16"] = "=INDEX(Ref_Scales!$A$5:$A$21,C15+1)"
    ws["B17"] = "Anchor short-term grade"
    ws["C17"] = ('=IF(OR(Inputs!$C$43="Superior",Inputs!$C$43="Strong"),'
                 'INDEX(Ref_Scales!$C$5:$C$21,C15+1),INDEX(Ref_Scales!$B$5:$B$21,C15+1))')
    ws["B18"] = "Anchor line  (a SCREEN, not the rating)"
    ws["C18"] = '=C16&" / "&C17'
    ws["C18"].font = mut

    # ----------------------------------------------------------------------- STAGE 2: overlay
    ws["B20"] = "STAGE 2 - COMMITTEE OVERLAY   (Step 17c: what the anchor cannot see)"
    ws["B20"].font = hdr
    ws["B20"].fill = navy_fill
    ws["C20"] = "Notches"
    ws["C20"].font = Font(bold=True, color="FFFFFF")
    ws["C20"].fill = navy_fill
    ws["D20"] = "Justification"
    ws["D20"].font = Font(bold=True, color="FFFFFF")
    ws["D20"].fill = navy_fill

    for i, (src_row, label, _hint) in enumerate(OVERLAY_ROWS):
        r = 21 + i
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value="=Inputs!C%d" % src_row)
        ws.cell(row=r, column=4, value="=Inputs!D%d" % src_row).alignment = wrap

    ws["B26"] = "Total overlay notches (+ improves)"
    ws["C26"] = "=SUM(C21:C25)"
    ws["C26"].font = sub
    ws["B27"] = "Index after overlay"
    ws["C27"] = "=MIN(16,MAX(0,C15-C26))"

    # ------------------------------------------------------------------ STAGE 3: the rating
    ws["B29"] = "STAGE 3 - ASSIGNED RATING   (the Committee's decision - what the rationale states)"
    ws["B29"].font = hdr
    ws["B29"].fill = teal_fill

    ws["B30"] = "Committee direct override index (blank = anchor + overlay)"
    # Renders "(none)" rather than empty when the Committee did not override: an
    # empty formula result is indistinguishable from an un-recalculated cell, and that
    # ambiguity made recalc_workbook.verify() report a false blank - which in turn
    # aborted the --all batch part-way through (found 02-Sep-2026).
    ws["C30"] = '=IF(Inputs!C66="","(none)",Inputs!C66)'
    ws["D30"] = "=Inputs!D66"
    ws["D30"].alignment = wrap
    ws["B31"] = "ASSIGNED INDEX"
    ws["C31"] = '=IF(Inputs!C66="",C27,Inputs!C66)'
    ws["C31"].font = sub
    ws["B32"] = "Assigned long-term grade"
    ws["C32"] = "=INDEX(Ref_Scales!$A$5:$A$21,C31+1)"
    ws["C32"].font = sub
    ws["B33"] = "Assigned short-term grade"
    ws["C33"] = ('=IF(OR(Inputs!$C$43="Superior",Inputs!$C$43="Strong"),'
                 'INDEX(Ref_Scales!$C$5:$C$21,C31+1),INDEX(Ref_Scales!$B$5:$B$21,C31+1))')

    ws["B34"] = "Divergence, anchor -> assigned (notches)"
    ws["C34"] = "=C15-C31"
    ws["D34"] = ('=IF(C34=0,"None - the anchor and the assigned rating agree.",'
                 '"COMPUTED, not typed. Stage 2 must carry a justification for every notch of '
                 'this. The rationale quotes THIS number.")')
    ws["D34"].font = mut
    ws["D34"].alignment = wrap

    ws["B36"] = "Assigned long-term line (agency house style)"
    ws["C36"] = ('=IF(Inputs!C5="ICRA","[ICRA]"&C32&IF(Inputs!C44<>"Not applicable",'
                 '" ("&Inputs!C44&")",""),IF(Inputs!C5="CRISIL","Crisil "&C32&'
                 'IF(Inputs!C44<>"Not applicable","/"&Inputs!C44,""),"CARE "&C32&'
                 'IF(Inputs!C44<>"Not applicable","; "&Inputs!C44,"")))')
    ws["B37"] = "Assigned short-term line"
    ws["C37"] = ('=IF(Inputs!C5="ICRA","[ICRA]"&C33,IF(Inputs!C5="CRISIL","Crisil "&C33,'
                 '"CARE "&C33))')
    ws["B38"] = "Assigned combined"
    ws["C38"] = '=C36&" / "&C37'

    # ------------------------------------------------------------------------- projections
    ws["B40"] = "Future earning potential (from Projections)"
    ws["B40"].font = sub
    ws["B41"] = "Rs crore / x"
    ws["C41"] = "FY+1"
    ws["D41"] = "FY+2"
    proj = [("Revenue", "C5", "D5"), ("EBITDA", "C6", "D6"), ("PAT", "C13", "D13"),
            ("Gross cash accruals", "C14", "D14"), ("Free cash flow", "C18", "D18")]
    for i, (label, c1, c2) in enumerate(proj):
        r = 42 + i
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value="=Projections!%s" % c1)
        ws.cell(row=r, column=4, value="=Projections!%s" % c2)
    ratio = [("Interest coverage (x)", "F14", "G14"), ("Gearing (x)", "F10", "G10"),
             ("Net debt", "F17", "G17")]
    for i, (label, c1, c2) in enumerate(ratio):
        r = 47 + i
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value="=Ratios!%s" % c1)
        ws.cell(row=r, column=4, value="=Ratios!%s" % c2)

    # ----------------------------------------------------------------------------- footer
    ws["B51"] = ("HOW TO READ THIS SHEET. Stage 1 is a SCREEN computed from the audited "
                 "financials alone - it is deliberately blind to franchise, parentage and "
                 "sector accounting conventions, so that anything those factors add has to be "
                 "argued rather than absorbed silently. Stage 2 is where that argument is "
                 "recorded, one justified notch at a time. Stage 3 is the Committee's "
                 "decision and is the ONLY grade on this sheet that is a rating.")
    ws["B51"].alignment = wrap
    ws["B51"].font = mut
    ws["B53"] = ("This workbook is an analytical exercise output, not an agency publication. "
                 "The rating decision belongs to the Rating Committee (Step 17c), never to "
                 "this sheet.")
    ws["B53"].font = mut
    ws["B53"].alignment = wrap

    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 62
    ws.column_dimensions["F"].width = 18


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="print the layout, change nothing")
    ap.add_argument("--file", default=str(TEMPLATE))
    ap.add_argument("--no-backup", action="store_true",
                    help="skip the sibling .pre_restructure backup (case files are backed "
                         "up out-of-tree instead, so the case folder stays clean)")
    a = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                                     # noqa: BLE001
        pass

    from openpyxl import load_workbook
    path = Path(a.file)

    if a.check:
        wb = load_workbook(path)
        ws = wb["Rating"]
        for r in range(1, ws.max_row + 1):
            cells = []
            for c in range(2, 5):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    cells.append("%s%d=%s" % ("BCD"[c - 2], r, str(v)[:70]))
            if cells:
                print(" | ".join(cells))
        return 0

    backup = path.with_suffix(".pre_restructure.xlsx")
    if not a.no_backup and not backup.exists():
        shutil.copyfile(path, backup)
        print("backed up  -> %s" % backup.name)

    wb = load_workbook(path)
    build_inputs_block(wb["Inputs"])
    build_rating_sheet(wb["Rating"])
    wb.save(path)
    print("restructured %s" % path.name)
    print("  Inputs!B58:D67  committee-overlay inputs + hard-cap disapplication")
    print("  Rating sheet    Stage 1 anchor (10-18) -> Stage 2 overlay (20-27) -> Stage 3 (29-38)")
    print("  headline        C3/C4/C5 now echo the ASSIGNED rating, not the anchor")
    return 0


if __name__ == "__main__":
    sys.exit(main())
